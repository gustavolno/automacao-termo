"""Envio dos boletos do acordo por e-mail pelo webmail da Locaweb (Roundcube).

Fluxo por cliente: Criar email (o corpo/assinatura já vem pronto) -> Para -> Assunto ->
anexa os boletos do mês vigente em diante (lotes de até 20 arquivos) -> Enviar.
"""
import io
import json
import os
import re
import sys
import time
from datetime import date
from pathlib import Path
from typing import Callable, Optional

try:
    from baixar_boletos.config import (WEBMAIL_URL, EMAILS_ENVIADOS_FILE, normalizar_texto)
    from baixar_boletos import config as _cfg
except ImportError:
    from config import (WEBMAIL_URL, EMAILS_ENVIADOS_FILE, normalizar_texto)
    import config as _cfg

ASSUNTO = "Boletos acordo Geap Saúde/Aldrigues Cândido Advocacia - {nome}"
MAX_ARQUIVOS_POR_UPLOAD = 20
EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
BOLETO_RE = re.compile(r"BOLETO\s*-\s*(?:ENTRADA\s*-\s*)?(\d{4})-(\d{2})", re.I)
NOME_LISTA = "LISTA DE ENVIO DE BOLETOS"


MESES_EXT = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho",
             "agosto", "setembro", "outubro", "novembro", "dezembro"]
DOMINIO_ESCRITORIO = "aldriguescandido.adv.br"


def rotulo_mes(mes: tuple) -> str:
    """(2026, 10) -> 'outubro/2026'."""
    return f"{MESES_EXT[mes[1] - 1]}/{mes[0]}"


def e_do_escritorio(email: str) -> bool:
    """E-mail do próprio escritório, inclusive com erro de leitura do OCR/termo
    (aldriquescandido.adyv.br, aldnguescand1do...)."""
    import difflib
    dominio = normalizar_texto(email.split("@")[-1]).lower()
    # parecido com o domínio do escritório E começando como ele ("aldrigues..."), para não pegar domínios como candido.adv.br
    return dominio.startswith("ald") and difflib.SequenceMatcher(None, dominio, DOMINIO_ESCRITORIO).ratio() >= 0.7


def corrigir_email(email: str) -> tuple:
    """(e-mail corrigido, aviso). Só conserta erros de digitação óbvios no fim do endereço."""
    e = email.strip().strip(".").lower()
    for errado, certo in ((".comm", ".com"), (".cam", ".com"), (".con", ".com"), (".br.br", ".br")):
        if e.endswith(errado):
            return e[:-len(errado)] + certo, f"corrigido de {email}"
    return e, ""

def limpar_nome(nome: str) -> str:
    """'ELIZANE ... MÁXIMO- (SILVA PIMENTEL)' -> 'ELIZANE ... MÁXIMO'."""
    nome = re.sub(r"\(.*?\)", "", str(nome or ""))
    return " ".join(nome.split()).strip(" -").upper()


def mes_do_boleto(arquivo) -> Optional[tuple]:
    """(ano, mês) pelo nome 'BOLETO - [ENTRADA - ]AAAA-MM (...)'."""
    m = BOLETO_RE.search(Path(arquivo).name)
    return (int(m.group(1)), int(m.group(2))) if m else None


def boletos_a_enviar(pasta: Path, hoje: Optional[date] = None) -> list:
    """PDFs de boleto com vencimento do mês vigente em diante (a partir do dia
    DIA_VIRADA_MES, do mês seguinte em diante), em ordem cronológica."""
    atual = _cfg.mes_inicial_boletos(hoje)
    itens = []
    for f in Path(pasta).glob("*.pdf"):
        m = BOLETO_RE.search(f.name)
        if m and (int(m.group(1)), int(m.group(2))) >= atual:
            itens.append(((int(m.group(1)), int(m.group(2)), "ENTRADA" not in f.name.upper()), f))
    return [f for _, f in sorted(itens)]


def _pastas_lista() -> list:
    base = Path(__file__).resolve().parent
    pastas = [Path(sys.executable).parent, base, base.parent / "dist", base.parent]
    if os.getenv("LISTA_EMAILS_DIR"):
        pastas.insert(0, Path(os.getenv("LISTA_EMAILS_DIR")))
    return pastas


def localizar_lista_emails() -> Optional[Path]:
    """A lista de envio mais recente; prefere a versão (FINAL)."""
    achados = []
    for p in _pastas_lista():
        if p.exists():
            achados += [f for f in p.glob(f"{NOME_LISTA}*.xlsx") if not f.name.startswith("~$")]
    if not achados:
        return None
    return max(achados, key=lambda f: ("(FINAL)" in f.name.upper(), f.stat().st_mtime))


def carregar_lista_emails(caminho: Optional[Path] = None) -> dict:
    """{NOME NORMALIZADO: {"nome", "emails", "pasta"}} a partir da lista de envio."""
    import openpyxl
    try:
        from baixar_boletos.acordo_parser import read_locked_file_bytes
    except ImportError:
        from acordo_parser import read_locked_file_bytes
    caminho = caminho or localizar_lista_emails()
    if not caminho:
        return {}
    ws = openpyxl.load_workbook(io.BytesIO(read_locked_file_bytes(str(caminho))), data_only=True).active
    cab = {normalizar_texto(str(ws.cell(1, c).value or "")): c for c in range(1, ws.max_column + 1)}
    c_nome, c_mail = cab.get("CLIENTE", 1), cab.get("E-MAIL", 4)
    c_pasta = cab.get("PASTA DOS BOLETOS")
    lista = {}
    for r in range(2, ws.max_row + 1):
        nome = ws.cell(r, c_nome).value
        if not nome:
            continue
        lista[normalizar_texto(limpar_nome(nome))] = {
            "nome": limpar_nome(nome),
            "emails": list(dict.fromkeys(EMAIL_RE.findall(str(ws.cell(r, c_mail).value or "")))),
            "pasta": str(ws.cell(r, c_pasta).value or "") if c_pasta else "",
        }
    return lista


# E-mails do escritório e de tribunais/CEJUSC aparecem nos termos e atas, mas não são do cliente
EMAILS_INSTITUCIONAIS = ("aldriguescandido", "jus.br", "cejusc", "cesrei", "mp.br", "defensoria", "geap")


def buscar_email_no_termo(nome: str) -> dict:
    """Procura o e-mail do cliente no termo de acordo (PDF com texto) da pasta dele no Dropbox.
    Devolve {"emails": [...], "arquivo": nome do PDF lido, "motivo": explicação quando não achou}."""
    try:
        import pymupdf
        from baixar_boletos.acordo_parser import AcordoParser
    except ImportError:
        import pymupdf
        from acordo_parser import AcordoParser
    p = AcordoParser()
    pastas = [x for x in (p._buscar_pasta_em_diretorio(b, nome) for b in
                          (p.dropbox_base_padrao, p.dropbox_base_juridico, p.dropbox_base_fechados)) if x]
    if not pastas:
        return {"emails": [], "arquivo": None, "motivo": "pasta do cliente não encontrada"}

    def pdfs(chaves):
        return [f for pz in pastas for f in Path(pz).iterdir()
                if f.is_file() and f.suffix.lower() == ".pdf" and any(k in normalizar_texto(f.name) for k in chaves)]
    termos = pdfs(("TERMO",)) or pdfs(("ACORDO", "ATA", "AUDIENCIA"))
    if not termos:
        return {"emails": [], "arquivo": None, "motivo": "nenhum termo de acordo na pasta"}

    escaneados = []
    for f in sorted(termos, key=lambda x: x.stat().st_mtime, reverse=True):  # o mais recente primeiro
        try:
            with pymupdf.open(f) as doc:
                texto = " ".join("".join(pg.get_text() for pg in doc).split())
        except Exception:
            continue
        if not texto.strip():
            escaneados.append(f.name)
            continue
        achados = [e.strip(".") for e in EMAIL_RE.findall(texto)
                   if not any(k in e.lower() for k in EMAILS_INSTITUCIONAIS)]
        if achados:
            return {"emails": list(dict.fromkeys(achados)), "arquivo": f.name, "motivo": ""}
    if escaneados:
        return {"emails": [], "arquivo": escaneados[0], "motivo": "termo escaneado (imagem), não dá para ler o e-mail"}
    return {"emails": [], "arquivo": termos[0].name, "motivo": "o termo não tem e-mail do cliente"}


_layna_cache = {}


def _carregar_layna() -> tuple:
    """(por_cpf, por_nome, arquivo) com os e-mails das abas da planilha da LAYNA
    ('LAYNA- CONTROLE DE PAGAMENTO - MENSAL GEAP.xlsx'), procurada junto da lista de envio."""
    achados = []
    for p in _pastas_lista():
        if p.exists():
            achados += [f for f in p.glob("*LAYNA*CONTROLE*PAGAMENTO*.xlsx") if not f.name.startswith("~$")]
    if not achados:
        return {}, {}, None
    arq = max(achados, key=lambda f: f.stat().st_mtime)
    chave = (str(arq), arq.stat().st_mtime)
    if _layna_cache.get("chave") == chave:
        return _layna_cache["por_cpf"], _layna_cache["por_nome"], arq
    import openpyxl
    try:
        from baixar_boletos.acordo_parser import read_locked_file_bytes
    except ImportError:
        from acordo_parser import read_locked_file_bytes
    wb = openpyxl.load_workbook(io.BytesIO(read_locked_file_bytes(str(arq))), data_only=True)
    por_cpf, por_nome = {}, {}
    for ws in wb.worksheets:
        cab = {normalizar_texto(str(ws.cell(1, c).value or "")): c for c in range(1, ws.max_column + 1)}
        c_mail = next((c for h, c in cab.items() if h in ("EMAIL", "E-MAIL")), None)
        c_nome = next((c for h, c in cab.items() if h in ("BENEFICIARIO", "CLIENTE", "NOME")), None)
        c_cpf = cab.get("CPF")
        if not c_mail or not c_nome:
            continue
        for r in range(2, ws.max_row + 1):
            mails = EMAIL_RE.findall(str(ws.cell(r, c_mail).value or ""))
            if not mails:
                continue
            cpf = re.sub(r"\D", "", str(ws.cell(r, c_cpf).value or "")) if c_cpf else ""
            nome = normalizar_texto(limpar_nome(ws.cell(r, c_nome).value))
            for m in mails:
                if cpf:
                    por_cpf.setdefault(cpf.lstrip("0"), []).append(m)
                if nome:
                    por_nome.setdefault(nome, []).append(m)
    _layna_cache.update(chave=chave, por_cpf=por_cpf, por_nome=por_nome)
    return por_cpf, por_nome, arq


def buscar_email_layna(nome: str, cpf: str = "") -> dict:
    """E-mail do cliente na planilha da LAYNA, pelo CPF (preferência) ou pelo nome."""
    try:
        por_cpf, por_nome, arq = _carregar_layna()
    except Exception as e:
        return {"emails": [], "arquivo": None, "motivo": f"erro ao ler a planilha da LAYNA: {e}"}
    if not arq:
        return {"emails": [], "arquivo": None, "motivo": "planilha da LAYNA não encontrada"}
    cpf = re.sub(r"\D", "", cpf or "").lstrip("0")
    emails, por = [], ""
    if cpf and por_cpf.get(cpf):
        emails, por = por_cpf[cpf], "pelo CPF"
    elif por_nome.get(normalizar_texto(limpar_nome(nome))):
        emails, por = por_nome[normalizar_texto(limpar_nome(nome))], "SÓ PELO NOME (confira se é a mesma pessoa)"
    emails = [e for e in dict.fromkeys(emails) if not any(k in e.lower() for k in EMAILS_INSTITUCIONAIS)]
    return {"emails": emails, "arquivo": arq.name, "por": por,
            "motivo": "" if emails else "cliente sem e-mail na planilha da LAYNA"}


def buscar_email_cliente(nome: str, lista: Optional[dict] = None) -> dict:
    """Procura o e-mail na ordem: lista de envio -> termo de acordo -> planilha da LAYNA.
    Devolve {"emails", "fonte", "pasta" (da lista), "motivos" (onde não achou e por quê)}."""
    item = (lista or {}).get(normalizar_texto(limpar_nome(nome)), {})
    if item.get("emails"):
        return {"emails": item["emails"], "fonte": "lista de envio", "pasta": item.get("pasta", ""), "motivos": []}
    motivos = ["lista de envio: cliente não está na lista" if not item else "lista de envio: sem e-mail"]
    try:
        termo = buscar_email_no_termo(nome)
    except Exception as e:
        termo = {"emails": [], "arquivo": None, "motivo": f"erro ao ler: {e}"}
    if termo["emails"]:
        return {"emails": termo["emails"], "fonte": f"termo ({termo['arquivo']})", "pasta": item.get("pasta", ""),
                "motivos": motivos}
    motivos.append(f"termo: {termo['motivo']}")
    cpf = ""
    try:
        try:
            from baixar_boletos.acordo_parser import AcordoParser
        except ImportError:
            from acordo_parser import AcordoParser
        cpf = AcordoParser().carregar_dados_cliente(nome).get("cpf", "")
    except Exception:
        pass
    layna = buscar_email_layna(nome, cpf)
    if layna["emails"]:
        return {"emails": layna["emails"], "fonte": f"planilha da LAYNA, achado {layna['por']}",
                "pasta": item.get("pasta", ""), "motivos": motivos}
    motivos.append(f"LAYNA: {layna['motivo']}")
    return {"emails": [], "fonte": "", "pasta": item.get("pasta", ""), "motivos": motivos}


def localizar_pasta_boletos(nome: str, pasta_lista: str = "") -> Optional[Path]:
    """Pasta BOLETOS DO ACORDO do cliente (sem criar pastas novas)."""
    if pasta_lista and Path(pasta_lista).exists():
        return Path(pasta_lista)
    try:
        from baixar_boletos.acordo_parser import AcordoParser
    except ImportError:
        from acordo_parser import AcordoParser
    p = AcordoParser()
    for base in (p.dropbox_base_juridico, p.dropbox_base_padrao, p.dropbox_base_fechados):
        pasta = p._buscar_pasta_em_diretorio(base, nome)
        if pasta and (pasta / "BOLETOS DO ACORDO").exists():
            return pasta / "BOLETOS DO ACORDO"
    return None


def carregar_historico() -> dict:
    """Histórico de envios. Arquivo ausente = nunca enviou nada. Arquivo que EXISTE mas não pôde ser lido
    (travado pelo OneDrive, truncado) levanta erro: devolver {} faria a Central achar que ninguém recebeu
    nada e reenviar para todos, e o próximo registro apagaria o histórico antigo."""
    arq = Path(EMAILS_ENVIADOS_FILE)
    if not arq.exists():
        return {}
    try:
        return json.loads(arq.read_text(encoding="utf-8"))
    except Exception as e:
        raise RuntimeError(f"O histórico de e-mails enviados ({arq.name}) existe mas não pôde ser lido ({e}). "
                           "Nada será enviado até isso ser resolvido, para não repetir e-mails.")


def ja_enviado(nome: str, arquivos: list) -> Optional[dict]:
    """Último envio deste cliente que já continha todos estes arquivos (None se não houve)."""
    nomes = {Path(a).name for a in arquivos}
    for envio in reversed(carregar_historico().get(normalizar_texto(nome), [])):
        if nomes and nomes <= set(envio.get("arquivos", [])):
            return envio
    return None


def ultimo_envio_do_mes(nome: str, mes: tuple) -> Optional[dict]:
    """O envio mais recente deste cliente que levou o boleto do mês (ano, mês), ou None."""
    for envio in reversed(carregar_historico().get(normalizar_texto(nome), [])):
        if any(mes_do_boleto(a) == mes for a in envio.get("arquivos", [])):
            return envio
    return None


def mes_ja_enviado(nome: str, mes: tuple) -> bool:
    """Algum e-mail já enviado a este cliente levou o boleto do mês (ano, mês)?"""
    return ultimo_envio_do_mes(nome, mes) is not None

def registrar_envio(nome: str, emails: list, arquivos: list):
    hist = carregar_historico()
    hist.setdefault(normalizar_texto(nome), []).append({
        "data": time.strftime("%Y-%m-%d %H:%M"), "nome": nome, "para": emails,
        "arquivos": [Path(a).name for a in arquivos],
    })
    arq = Path(EMAILS_ENVIADOS_FILE)
    tmp = arq.with_suffix(".tmp")
    tmp.write_text(json.dumps(hist, ensure_ascii=False, indent=1), encoding="utf-8")
    os.replace(tmp, arq)   # troca atômica: uma queda no meio não deixa o histórico pela metade


class EmailService:
    def __init__(self, log_callback: Callable = print, cancelado: Callable[[], bool] = lambda: False):
        self.log = log_callback
        self.cancelado = cancelado
        self.playwright = self.browser = self.page = None
        self._alertas = []

    # ── sessão ──
    def abrir_sessao(self, visivel: bool = False):
        if self.page and not self.page.is_closed():
            return
        usuario, senha = _cfg.WEBMAIL_USER, os.getenv("WEBMAIL_PASS", _cfg.WEBMAIL_PASS)
        if not senha:
            raise RuntimeError("Senha do webmail não configurada (⚙️ Meu Login > Webmail).")
        from playwright.sync_api import sync_playwright
        self.playwright = sync_playwright().start()
        self.browser = self.playwright.chromium.launch(headless=not visivel)
        self.page = self.browser.new_page(viewport={"width": 1500, "height": 950})
        # alert()/confirm() do Roundcube (ex.: endereço inválido) viram erro em vez de travar
        self.page.on("dialog", self._ao_dialogo)
        self.log("Webmail: entrando...")
        self.page.goto(WEBMAIL_URL, wait_until="domcontentloaded")
        self.page.fill("#rcmloginuser", usuario)
        self.page.fill("#rcmloginpwd", senha)
        self.page.keyboard.press("Enter")
        try:
            self.page.wait_for_url("**_task=mail**", timeout=40000)
        except Exception:
            raise RuntimeError("Login no webmail falhou (confira usuário/senha).")
        self.page.wait_for_timeout(2000)
        self._fechar_modais()
        self.log("Webmail: login OK.")

    def _ao_dialogo(self, dialog):
        self._alertas.append(dialog.message)
        try:
            dialog.dismiss()
        except Exception:
            pass

    def _fechar_modais(self):
        """Fecha o aviso 'Cadastre um e-mail alternativo' (e similares) se aparecer."""
        modal = self.page.locator("#alternative_email_modal.ls-opened")
        if modal.count():
            try:
                modal.get_by_text("Fechar").first.click(timeout=3000)
                self.page.wait_for_timeout(800)
            except Exception:
                pass

    def fechar_sessao(self):
        # Cada passo isolado: se o navegador já foi fechado (ex.: janela fechada à mão), o erro
        # do close() não pode impedir o stop() — sem ele, a próxima sessão falha com
        # "Playwright Sync API inside the asyncio loop".
        for passo in (lambda: self.browser and self.browser.close(),
                      lambda: self.playwright and self.playwright.stop()):
            try:
                passo()
            except Exception:
                pass
        self.playwright = self.browser = self.page = None

    # ── e-mail ──
    def _abrir_compose(self):
        pg = self.page
        pg.goto(WEBMAIL_URL.rstrip("/") + "/?_task=mail&_action=compose", wait_until="domcontentloaded")
        pg.wait_for_selector("#compose-subject", timeout=30000)
        # espera o modelo pré-pronto (corpo + assinatura) carregar no editor
        for _ in range(20):
            if pg.evaluate("window.tinymce && tinymce.activeEditor && tinymce.activeEditor.initialized "
                           "&& tinymce.activeEditor.getContent({format:'text'}).includes('Prezado')"):
                break
            pg.wait_for_timeout(500)
        else:
            self.log("Aviso: o texto padrão do e-mail não apareceu no editor.", "WARNING")
        self._fechar_modais()

    def _anexar(self, arquivos: list):
        pg = self.page
        for i in range(0, len(arquivos), MAX_ARQUIVOS_POR_UPLOAD):
            lote = [str(a) for a in arquivos[i:i + MAX_ARQUIVOS_POR_UPLOAD]]
            antes = pg.locator("#attachment-list li").count()
            try:
                pg.locator("#fileuploadbtn").click(timeout=5000)
                pg.wait_for_timeout(600)
            except Exception:
                pass
            pg.set_input_files("#uploadformInput", lote)
            pg.evaluate("rcmail.upload_file(document.querySelector('#uploadformInput').form, 'upload')")
            alvo = antes + len(lote)
            for _ in range(120):
                n = pg.locator("#attachment-list li").count()
                enviando = pg.locator("#attachment-list li.uploading").count()
                if n >= alvo and not enviando:
                    break
                pg.wait_for_timeout(500)
            else:
                raise RuntimeError(f"Anexos não terminaram de subir ({pg.locator('#attachment-list li').count()}/{alvo}).")
            # fecha o diálogo "Anexar um arquivo" se continuou aberto
            pg.evaluate("""() => { for (const b of document.querySelectorAll('.ui-dialog button, .modal button, button'))
                                     if (b.offsetParent && /^\\s*Cancelar\\s*$/.test(b.innerText)) { b.click(); return; } }""")
            pg.wait_for_timeout(400)
        nomes = pg.evaluate("[...document.querySelectorAll('#attachment-list li .attachment-name')].map(e => e.innerText.trim())")
        faltando = [Path(a).name for a in arquivos if Path(a).name not in nomes]
        if faltando:
            raise RuntimeError(f"Anexos ausentes no e-mail: {faltando}")

    def descartar(self):
        """Sai do compose sem enviar e apaga o rascunho, se o webmail tiver salvo um."""
        try:
            self.page.evaluate("rcmail.command('locaweb_save_draft.removeMessage')")
            self.page.wait_for_timeout(1500)
        except Exception:
            pass
        try:
            self.page.goto(WEBMAIL_URL.rstrip("/") + "/?_task=mail&_mbox=INBOX", wait_until="domcontentloaded")
        except Exception:
            pass

    def preparar_email(self, nome: str, emails: list, arquivos: list):
        """Abre um e-mail novo e preenche Para, Assunto e anexos (não envia)."""
        self._alertas.clear()
        self._abrir_compose()
        pg = self.page
        pg.locator("#_to").fill(", ".join(emails))
        pg.locator("#compose-subject").fill(ASSUNTO.format(nome=nome))
        self._anexar(arquivos)

    def enviar(self) -> bool:
        pg = self.page
        self._alertas.clear()
        pg.locator("[data-testid='messages.compose.send']").click(timeout=10000)
        for _ in range(120):
            if self._alertas:
                raise RuntimeError(f"Webmail recusou o envio: {self._alertas[-1]}")
            if "_action=compose" not in pg.url:
                return True
            pg.wait_for_timeout(500)
        raise RuntimeError("Sem confirmação de envio após 60s: o e-mail PODE ter saído. Confira a pasta Enviados do webmail antes de tentar de novo.")

    def enviar_mes(self, nome: str, emails: list, arquivos: list, mes: tuple,
                   confirmar: Optional[Callable[[str], str]] = None, visivel: bool = False) -> dict:
        """Envia SÓ o(s) boleto(s) do mês `mes` ((ano, mês)), com o texto do e-mail ajustado para
        esse mês. Se o texto padrão do webmail mudou e algum trecho não for achado, não envia."""
        n, rot = len(arquivos), rotulo_mes(mes)
        if not arquivos:
            return {"status": "sem_boletos", "msg": "nenhum boleto do mês"}

        def ent(t):   # acentos como entidade HTML, como o editor do webmail devolve
            return t.encode("ascii", "xmlcharrefreplace").decode()
        antigo = "<span>Informamos que os </span><strong><span>boletos referentes ao seu acordo"
        if n == 1:
            trocas = [
                (antigo, f"<span>Informamos que o </span><strong><span>boleto do m&ecirc;s de {ent(rot)} referente ao seu acordo"),
                ("<span>&nbsp;est&atilde;o dispon&iacute;veis e seguem anexados a este e-mail.</span>",
                 "<span>&nbsp;est&aacute; dispon&iacute;vel e segue anexado a este e-mail.</span>"),
                ("que observe atentamente as </span><strong><span>datas de vencimento</span></strong><span> de cada boleto e realize os pagamentos dentro dos respectivos prazos,",
                 "que observe atentamente a </span><strong><span>data de vencimento</span></strong><span> do boleto e realize o pagamento dentro do prazo,"),
                ("Ap&oacute;s a realiza&ccedil;&atilde;o de cada pagamento,", "Ap&oacute;s a realiza&ccedil;&atilde;o do pagamento,"),
            ]
        else:   # entrada + parcela no mesmo mês
            trocas = [(antigo, f"<span>Informamos que os </span><strong><span>boletos do m&ecirc;s de {ent(rot)} (entrada e parcela) referentes ao seu acordo")]

        self.abrir_sessao(visivel=visivel)
        self.log(f"{nome}: {n} boleto(s) de {rot} -> {', '.join(emails)}")
        self.preparar_email(nome, emails, arquivos)
        html = self.page.evaluate("tinymce.activeEditor.getContent()")
        for a, b in trocas:
            if html.count(a) != 1:
                self.descartar()
                raise RuntimeError(f"o texto padrão do e-mail mudou (trecho não encontrado: {a[:50]}...); nada enviado")
            html = html.replace(a, b)
        self.page.evaluate("h => tinymce.activeEditor.setContent(h)", html)
        if rot not in self.page.evaluate("tinymce.activeEditor.getContent({format:'text'})"):
            self.descartar()
            raise RuntimeError("o texto do e-mail não ficou ajustado para o mês; nada enviado")
        if self.cancelado():
            self.descartar()
            return {"status": "parado"}
        if confirmar:
            resp = confirmar(f"{nome}\nPara: {', '.join(emails)}\nBoleto(s) de {rot}: {n}\n\nEnviar este e-mail?")
            if resp != "enviar":
                self.descartar()
                return {"status": "parado" if resp == "parar" else "pulado"}
        self.enviar()
        registrar_envio(nome, emails, arquivos)
        self.log(f"{nome}: e-mail ENVIADO com {n} boleto(s) de {rot}.", "SUCCESS")
        return {"status": "enviado", "arquivos": n}

    def enviar_boletos_cliente(self, nome: str, emails: list, pasta: Path,
                               confirmar: Optional[Callable[[str], str]] = None, visivel: bool = False) -> dict:
        """confirmar(msg) -> 'enviar' | 'pular' | 'parar'. Sem confirmar, envia direto."""
        arquivos = boletos_a_enviar(pasta)
        if not arquivos:
            return {"status": "sem_boletos", "msg": "nenhum boleto do mês vigente em diante na pasta"}
        self.abrir_sessao(visivel=visivel)
        self.log(f"{nome}: {len(arquivos)} boletos ({arquivos[0].name} ... {arquivos[-1].name}) -> {', '.join(emails)}")
        self.preparar_email(nome, emails, arquivos)
        if self.cancelado():
            self.descartar()
            return {"status": "parado"}
        if confirmar:
            resp = confirmar(f"{nome}\nPara: {', '.join(emails)}\nAnexos: {len(arquivos)} boletos\n\nEnviar este e-mail?")
            if resp != "enviar":
                self.descartar()
                return {"status": "parado" if resp == "parar" else "pulado"}
        self.enviar()
        registrar_envio(nome, emails, arquivos)
        self.log(f"{nome}: e-mail ENVIADO com {len(arquivos)} boletos.", "SUCCESS")
        return {"status": "enviado", "arquivos": len(arquivos)}
