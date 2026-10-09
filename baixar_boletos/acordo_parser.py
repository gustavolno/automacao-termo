try:
    from baixar_boletos.config import obter_pasta_boletos_cliente
except ImportError:
    from config import obter_pasta_boletos_cliente

"""
acordo_parser.py
Módulo para leitura e validação dos dados de acordos (apenas PARCELADOS)
cruzando informações da planilha com termos e atas de audiência em PDF no Dropbox.
"""

import os
import sys
import re
import unicodedata
from datetime import datetime, date
from pathlib import Path
from typing import Optional, Dict, Any, Union

import io
import ctypes
from ctypes import wintypes

import openpyxl
import pymupdf

try:
    from baixar_boletos.config import PLANILHA_ONLINE_FILE, PLANILHA_DROPBOX_FILE
except ImportError:
    from config import PLANILHA_ONLINE_FILE, PLANILHA_DROPBOX_FILE


def read_locked_file_bytes(filepath: str) -> bytes:
    """Lê bytes de um arquivo mesmo com lock exclusivo do Excel no Windows usando CreateFileW."""
    if os.name != 'nt':
        with open(filepath, 'rb') as f:
            return f.read()

    CreateFileW = ctypes.windll.kernel32.CreateFileW
    CreateFileW.argtypes = [
        wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD,
        wintypes.LPVOID, wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE
    ]
    CreateFileW.restype = wintypes.HANDLE

    ReadFile = ctypes.windll.kernel32.ReadFile
    ReadFile.argtypes = [
        wintypes.HANDLE, wintypes.LPVOID, wintypes.DWORD,
        ctypes.POINTER(wintypes.DWORD), wintypes.LPVOID
    ]
    ReadFile.restype = wintypes.BOOL

    GetFileSizeEx = ctypes.windll.kernel32.GetFileSizeEx
    GetFileSizeEx.argtypes = [wintypes.HANDLE, ctypes.POINTER(ctypes.c_int64)]
    GetFileSizeEx.restype = wintypes.BOOL

    CloseHandle = ctypes.windll.kernel32.CloseHandle
    CloseHandle.argtypes = [wintypes.HANDLE]
    CloseHandle.restype = wintypes.BOOL

    handle = CreateFileW(
        filepath,
        0x80000000,  # GENERIC_READ
        0x00000001 | 0x00000002 | 0x00000004,  # FILE_SHARE_READ | FILE_SHARE_WRITE | FILE_SHARE_DELETE
        None,
        3,  # OPEN_EXISTING
        0x80,  # FILE_ATTRIBUTE_NORMAL
        None
    )

    if handle == wintypes.HANDLE(-1).value:
        err = ctypes.GetLastError()
        raise PermissionError(f'CreateFileW falhou com código: {err}')

    try:
        size = ctypes.c_int64()
        if not GetFileSizeEx(handle, ctypes.byref(size)):
            raise IOError('Não foi possível obter tamanho do arquivo.')

        file_size = size.value
        buffer = ctypes.create_string_buffer(file_size)
        bytes_read = wintypes.DWORD()

        if not ReadFile(handle, buffer, file_size, ctypes.byref(bytes_read), None):
            err = ctypes.GetLastError()
            raise IOError(f'ReadFile falhou com código: {err}')

        return buffer.raw[:bytes_read.value]
    finally:
        CloseHandle(handle)


def carregar_planilha_segura(caminho: Union[str, Path]) -> openpyxl.Workbook:
    """Carrega uma planilha do Excel de forma segura, contornando bloqueios do Excel no Windows."""
    c = Path(caminho)
    try:
        return openpyxl.load_workbook(str(c), data_only=True)
    except Exception:
        data = read_locked_file_bytes(str(c.resolve()))
        return openpyxl.load_workbook(io.BytesIO(data), data_only=True)


def normalizar_texto(texto: str) -> str:
    """Remove acentos e espaços extras para comparações seguras."""
    if not texto:
        return ""
    texto = unicodedata.normalize("NFKD", str(texto)).encode("ASCII", "ignore").decode("utf-8")
    return re.sub(r"\s+", " ", texto).strip().upper()


def limpar_processo(proc_str: Any) -> str:
    """Extrai número CNJ no formato 0000000-00.0000.8.00.0000."""
    if not proc_str:
        return ""
    s = str(proc_str).strip()
    m = re.search(r"(\d{7}-\d{2}\.\d{4}\.\d\.\d{2}\.\d{4})", s)
    if m:
        return m.group(1)
    m2 = re.search(r"(\d{7}\s*\d{2}\s*\d{4}\s*\d\s*\d{2}\s*\d{4})", s)
    if m2:
        return re.sub(r"\s+", ".", m2.group(1))
    return s.split()[0].strip()


def formatar_moeda_br(valor: Any) -> str:
    """Converte valores numéricos para string formatada em Real: 1.234,56."""
    if valor is None:
        return ""
    s = str(valor).strip()
    if not s or s in ("-", "None", "0", "0.0", "0,00"):
        return ""
    try:
        if "," in s and "." in s:
            s = s.replace(".", "").replace(",", ".")
        elif "," in s:
            s = s.replace(",", ".")
        f = float(s)
        if f <= 0:
            return ""
        partes = f"{f:,.2f}".split(".")
        inteiro = partes[0].replace(",", ".")
        decimal = partes[1]
        return f"{inteiro},{decimal}"
    except Exception:
        return s


def formatar_data_br(val: Any) -> str:
    """Converte datetime, date ou string para dd/mm/aaaa."""
    if val is None:
        return ""
    if isinstance(val, (datetime, date)):
        return val.strftime("%d/%m/%Y")
    s = str(val).strip()
    if not s or s in ("-", "None"):
        return ""
    if re.match(r"^\d{2}/\d{2}/\d{4}$", s):
        return s
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        ano, mes, dia = m.group(1), m.group(2), m.group(3)
        return f"{dia}/{mes}/{ano}"
    return s


def extrair_qtd_parcelas(val: Any) -> int:
    """Extrai o número inteiro de parcelas: ex '24x' -> 24."""
    if not val:
        return 0
    m = re.search(r"(\d+)", str(val))
    if m:
        return int(m.group(1))
    return 0


class AcordoParser:
    def __init__(self, caminho_planilha: Optional[str] = None):
        self.caminho_planilha = self._resolver_caminho_planilha(caminho_planilha)
        self._cached_ws = None

        # Resolução resiliente das pastas base do Dropbox
        usuario = os.environ.get("USERNAME", "Cob_Geap_05082026")
        user_profile = Path(os.environ.get("USERPROFILE", f"C:/Users/{usuario}"))
        dropbox_dir = user_profile / "Dropbox"
        if not dropbox_dir.exists():
            dropbox_dir = Path(r"C:\Users\Cob_Geap_05082026\Dropbox")

        admin_cob = dropbox_dir / "COBRANÇA NEGOCIAÇÃO" / "ADMINISTRATIVO- COBRANÇA"
        self.dropbox_base_padrao = admin_cob / "ACOMPANHAMENTO DE ACORDOS AC COBRANÇA"
        self.dropbox_base_juridico = admin_cob / "ACOMPANHAMENTO DE ACORDOS - JURÍDICO"
        self.dropbox_base_fechados = admin_cob / "ACOMPANHAMENTO DE ACORDOS FECHADOS"
        self.dropbox_base = self.dropbox_base_padrao

    def _obter_worksheet(self):
        """Retorna a worksheet ativa com cache em memória sem recursão."""
        if self._cached_ws is not None:
            return self._cached_ws
        if not self.caminho_planilha or not self.caminho_planilha.exists():
            return None
        try:
            wb = carregar_planilha_segura(self.caminho_planilha)
            ws = wb.active
            for s in wb.sheetnames:
                if "FECHADO" in s.upper() or "ENCONTRADO" in s.upper():
                    ws = wb[s]
                    break
            self._cached_ws = ws
            return self._cached_ws
        except Exception:
            return None

    def cliente_eh_juridico(self, nome_cliente: str) -> bool:
        """Retorna True se o acordo NÃO for 'FECHADO AC COBRANÇA'."""
        try:
            ws = self._obter_worksheet()
            if ws is None:
                return False
            norm_alvo = normalizar_texto(nome_cliente)
            col_nome = 3
            col_fechado = 6
            for c in range(1, min(ws.max_column + 1, 30)):
                v = ws.cell(1, c).value
                if v:
                    nv = normalizar_texto(str(v))
                    if "BENEFICIARIO" in nv:
                        col_nome = c
                    elif "FECHADO EM" in nv:
                        col_fechado = c
            for r in range(2, ws.max_row + 1):
                val_nome = ws.cell(r, col_nome).value
                if val_nome:
                    nv = normalizar_texto(str(val_nome))
                    if nv == norm_alvo or (len(norm_alvo) > 8 and (norm_alvo in nv or nv in norm_alvo)):
                        val_fechado = str(ws.cell(r, col_fechado).value or "").strip()
                        norm_fech = normalizar_texto(val_fechado)
                        return bool(norm_fech and "AC COBRANCA" not in norm_fech)
        except Exception:
            pass
        return False

    def cliente_fechado_em_audiencia(self, nome_cliente: str) -> bool:
        return self.cliente_eh_juridico(nome_cliente)

    def _buscar_pasta_em_diretorio(self, diretorio_base: Path, nome_cliente: str) -> Optional[Path]:
        """
        Busca a pasta do cliente com correspondência estrita e segura:
        1. Correspondência exata de nome.
        2. Correspondência completa de substring (ex: 'NOME DO CLIENTE - ACORDO').
        3. Correspondência precisa: OBRIGATÓRIO que o primeiro E o segundo nome confiram.
           Isso IMPEDE que 'MARIA CRISTINA' seja confundida com 'MARIA HELENA'.
        """
        if not diretorio_base.exists():
            return None

        norm_alvo = normalizar_texto(nome_cliente)

        # 1. Busca exata de pasta
        for p in diretorio_base.iterdir():
            if p.is_dir() and normalizar_texto(p.name) == norm_alvo:
                return p

        # 2. Nome completo contido no nome da pasta
        for p in diretorio_base.iterdir():
            if p.is_dir():
                p_norm = normalizar_texto(p.name)
                if norm_alvo in p_norm or (len(p_norm) > 12 and p_norm in norm_alvo):
                    return p

        # 3. Busca por palavras: exige que os DOIS primeiros nomes e o sobrenome confiram
        palavras = [w for w in norm_alvo.split() if len(w) > 2]
        if len(palavras) >= 2:
            primeiros_dois = f"{palavras[0]} {palavras[1]}"  # Ex: "MARIA CRISTINA"
            for p in diretorio_base.iterdir():
                if p.is_dir():
                    p_norm = normalizar_texto(p.name)
                    # O primeiro e segundo nome têm que estar presentes juntos na pasta
                    if primeiros_dois in p_norm:
                        if len(palavras) == 2 or palavras[-1] in p_norm:
                            return p

        return None

    def _resolver_caminho_planilha(self, custom_path: Optional[str] = None) -> Path:
        if custom_path and Path(custom_path).exists():
            return Path(custom_path)

        # Planilha OFICIAL: a cópia mais recente entre a baixada do OneDrive e a do Dropbox
        oficiais = [p for p in (PLANILHA_ONLINE_FILE, PLANILHA_DROPBOX_FILE) if p.exists()]
        if oficiais:
            return max(oficiais, key=lambda p: p.stat().st_mtime).resolve()

        candidatos = [
            Path("devedores_encontrados.xlsx"),
            Path("planilha_jaci.xlsx"),
            Path(r"C:/Users/Cob_Geap_05082026/OneDrive/Documentos/automacao-termo/devedores_encontrados.xlsx"),
            Path(r"C:/Users/Cob_Geap_05082026/Dropbox/COBRANÇA NEGOCIAÇÃO/GERADOR DE TERMO/planilha_jaci.xlsx"),
            Path(r"C:/Users/Cob_Geap_05082026/Dropbox/COBRANÇA NEGOCIAÇÃO/planilha_jaci.xlsx"),
            Path(r"C:/Users/Cob_Geap_05082026/OneDrive/Documentos/automacao-termo/planilha_jaci.xlsx"),
            Path(os.environ.get("LOCALAPPDATA", "")) / "BaixarBoletos" / "planilha_jaci.xlsx",
        ]
        for c in candidatos:
            if c.exists():
                return c.resolve()
        return Path("devedores_encontrados.xlsx")

    def obter_pasta_cliente_dropbox(self, nome_cliente: str, eh_audiencia: Optional[bool] = None, eh_juridico: Optional[bool] = None) -> Optional[Path]:
        """Busca a pasta do cliente em todas as bases do Dropbox na ordem de prioridade."""
        if eh_juridico is None:
            eh_juridico = eh_audiencia if eh_audiencia is not None else self.cliente_eh_juridico(nome_cliente)

        # Ordem de busca conforme a origem do acordo
        if eh_juridico:
            bases = [self.dropbox_base_juridico, self.dropbox_base_fechados, self.dropbox_base_padrao]
        else:
            bases = [self.dropbox_base_padrao, self.dropbox_base_fechados, self.dropbox_base_juridico]

        for base in bases:
            pasta = self._buscar_pasta_em_diretorio(base, nome_cliente)
            if pasta:
                return pasta

        return None

    def verificar_boletos_cliente(self, nome_cliente: str) -> Dict[str, Any]:
        """
        Verifica se a pasta do cliente no Dropbox (.../BOLETOS DO ACORDO)
        já contém todos os boletos correspondentes às parcelas da planilha,
        considerando também os comprovantes de parcelas já pagas (ex: COMP DE PAG).
        """
        dados = self.carregar_dados_cliente(nome_cliente)
        if not dados.get("sucesso") or not dados.get("eh_parcelado"):
            return {
                "completo": False,
                "motivo": dados.get("motivo_erro", "Não parcelado"),
                "encontrados": 0,
                "esperados": 0,
                "pasta": None,
                "dados": dados
            }

        qtd_esperada = dados.get("qtd_parcelas", 0)
        if qtd_esperada <= 0:
            return {
                "completo": False,
                "motivo": "Parcelas não definidas",
                "encontrados": 0,
                "esperados": 0,
                "pasta": None,
                "dados": dados
            }

        pasta_dropbox = dados.get("pasta_dropbox")
        pasta_boletos = None

        if pasta_dropbox:
            p_drop = Path(pasta_dropbox)
            # 1. Procura subpasta BOLETOS DO ACORDO ou equivalente
            candidata = p_drop / "BOLETOS DO ACORDO"
            if candidata.exists():
                pasta_boletos = candidata
            else:
                for sub in p_drop.iterdir():
                    if sub.is_dir() and "BOLETO" in normalizar_texto(sub.name):
                        pasta_boletos = sub
                        break

        if not pasta_boletos or not pasta_boletos.exists():
            eh_jur = dados.get("eh_juridico", False)
            pasta_boletos = obter_pasta_boletos_cliente(nome_cliente, eh_juridico=eh_jur)

        if not pasta_boletos or not pasta_boletos.exists():
            return {
                "completo": False,
                "motivo": "Pasta de boletos não existe",
                "encontrados": 0,
                "esperados": qtd_esperada,
                "pasta": str(pasta_boletos) if pasta_boletos else None,
                "dados": dados
            }

        # 1. Boletos emitidos na subpasta de boletos
        boletos = [f for f in pasta_boletos.iterdir() if f.is_file() and f.suffix.lower() == ".pdf"]
        qtd_boletos = len(boletos)

        # 2. Comprovantes de pagamento na pasta principal do cliente
        qtd_comprovantes = 0
        pasta_pai = pasta_boletos.parent if pasta_boletos.name.upper().startswith("BOLETO") else pasta_boletos
        if pasta_pai.exists():
            for f in pasta_pai.iterdir():
                if f.is_file() and f.suffix.lower() == ".pdf":
                    nome_f = f.name.upper()
                    if any(k in nome_f for k in ["COMP", "PAG", "RECIBO"]) and "TERMO" not in nome_f:
                        qtd_comprovantes += 1

        qtd_encontrada = max(qtd_boletos, qtd_boletos + qtd_comprovantes)
        completo = (qtd_boletos >= qtd_esperada) or ((qtd_boletos + qtd_comprovantes) >= qtd_esperada)

        # Regra do mês vigente: só importam as parcelas do mês atual em diante (meses que já
        # passaram não são baixados). Completo = todas essas parcelas já estão na pasta.
        if not completo and dados.get("primeiro_vencimento"):
            try:
                from baixar_boletos.boletos_faltantes import cronograma_parcelas
                from baixar_boletos.config import mes_inicial_boletos
            except ImportError:
                from boletos_faltantes import cronograma_parcelas
                from config import mes_inicial_boletos
            inicio = mes_inicial_boletos()
            esperados = {(d.year, d.month) for d in cronograma_parcelas(dados["primeiro_vencimento"], qtd_esperada)
                         if (d.year, d.month) >= inicio}
            na_pasta = set()
            for f in boletos:
                m = re.search(r"BOLETO\s*-\s*(?:ENTRADA\s*-\s*)?(\d{4})-(\d{2})", f.name, re.I)
                if m:
                    na_pasta.add((int(m.group(1)), int(m.group(2))))
            if esperados <= na_pasta:
                completo = True
                qtd_encontrada, qtd_esperada = len(esperados & na_pasta), len(esperados)

        return {
            "completo": completo,
            "motivo": "Boletos completos" if completo else "Faltam boletos",
            "encontrados": qtd_encontrada,
            "boletos_gerados": qtd_boletos,
            "comprovantes_pagos": qtd_comprovantes,
            "esperados": qtd_esperada,
            "pasta": str(pasta_boletos),
            "dados": dados
        }

    def carregar_dados_cliente(self, nome_cliente: str) -> Dict[str, Any]:
        """Lê os dados da planilha e valida com a pasta do cliente no Dropbox de forma exata."""
        resultado = {
            "sucesso": False,
            "motivo_erro": "",
            "eh_parcelado": False,
            "nome_cliente": nome_cliente.strip(),
            "processo": "",
            "cpf": "",
            "data_registro": datetime.now().strftime("%d/%m/%Y"),
            "valor_inicial": "",
            "valor_limite": "",
            "vinculo": "Titular",
            "tem_sentenca": False,
            "arquivo_sentenca": None,
            "data_homologacao": "",
            "valor_homologacao": "",
            "descricao_item": "Prestação de serviço de cobrança/ tecnológico",
            "forma_pagamento": "Boleto bancário",
            "boleto": "SICOOB- AC",
            "escritorio": "Unidade Brasilia",
            "tem_entrada": False,
            "data_entrada": "",
            "valor_entrada": "",
            "qtd_parcelas": 0,
            "valor_parcela": "",
            "primeiro_vencimento": "",
            "pasta_dropbox": None,
            "documentos_encontrados": [],
            "validacao_pdf_notas": []
        }

        ws = self._obter_worksheet()
        if ws is None:
            resultado["motivo_erro"] = f"Planilha não encontrada: {self.caminho_planilha}"
            return resultado

        # Mapear cabeçalhos
        header = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        col_map = {}
        for idx, h in enumerate(header, 1):
            if h:
                col_map[normalizar_texto(str(h))] = idx

        idx_nome = col_map.get("BENEFICIARIO", 3)
        idx_cpf = col_map.get("CPF", 4)
        idx_proc = col_map.get("PROCESSO", 5)
        idx_fechado_em = col_map.get("ACORDO FECHADO EM", 6)
        idx_cond = col_map.get("CONDICOES DO ACORDO", 7)
        idx_vcausa = col_map.get("VALOR DA CAUSA", 8)
        idx_vacordo = col_map.get("VALOR INTEGRAL DO ACORDO", 10)
        idx_dentrada = col_map.get("DATA DA ENTRADA /QUITACAO", 11)
        idx_ventrada = col_map.get("VALOR DA ENTRADA/ QUITACAO", 12)
        idx_parc = col_map.get("PARCELAS", 15)
        idx_vparc = col_map.get("VALOR DA PARCELA", 16)
        idx_dprim_venc = col_map.get("PAGAMENTO INICIAL", 19)

        norm_alvo = normalizar_texto(nome_cliente)
        linha_encontrada = None

        # 1ª tentativa: Correspondência EXATA
        for r in range(2, ws.max_row + 1):
            val_nome = ws.cell(r, idx_nome).value
            if val_nome and normalizar_texto(str(val_nome)) == norm_alvo:
                linha_encontrada = r
                break

        # 2ª tentativa: Substring se não encontrou exato
        if not linha_encontrada:
            for r in range(2, ws.max_row + 1):
                val_nome = ws.cell(r, idx_nome).value
                if val_nome and len(norm_alvo) > 8:
                    nv = normalizar_texto(str(val_nome))
                    if norm_alvo in nv or nv in norm_alvo:
                        linha_encontrada = r
                        break

        if not linha_encontrada:
            resultado["motivo_erro"] = f"Cliente '{nome_cliente}' não localizado na planilha."
            return resultado

        # Extrair dados da linha
        val_fechado_em = str(ws.cell(linha_encontrada, idx_fechado_em).value or "").strip()
        norm_fech = normalizar_texto(val_fechado_em)
        eh_juridico = bool(norm_fech and "AC COBRANCA" not in norm_fech)
        resultado["fechado_em"] = val_fechado_em
        resultado["eh_audiencia"] = "AUDI" in norm_fech
        resultado["eh_juridico"] = eh_juridico

        condicao = str(ws.cell(linha_encontrada, idx_cond).value or "").strip().upper()
        if "PARCEL" not in condicao:
            resultado["eh_parcelado"] = False
            resultado["motivo_erro"] = f"Condição do acordo é '{condicao}' (Não é PARCELADO)."
            return resultado

        resultado["eh_parcelado"] = True
        resultado["cpf"] = str(ws.cell(linha_encontrada, idx_cpf).value or "").strip()
        resultado["processo"] = limpar_processo(ws.cell(linha_encontrada, idx_proc).value)
        resultado["valor_inicial"] = formatar_moeda_br(ws.cell(linha_encontrada, idx_vcausa).value)
        valor_acordo_str = formatar_moeda_br(ws.cell(linha_encontrada, idx_vacordo).value)
        resultado["valor_limite"] = valor_acordo_str
        resultado["valor_homologacao"] = valor_acordo_str

        # Entrada
        data_ent = formatar_data_br(ws.cell(linha_encontrada, idx_dentrada).value)
        val_ent = formatar_moeda_br(ws.cell(linha_encontrada, idx_ventrada).value)
        if val_ent and val_ent not in ("0,00", ""):
            resultado["tem_entrada"] = True
            resultado["data_entrada"] = data_ent
            resultado["valor_entrada"] = val_ent

        # Parcelas
        resultado["qtd_parcelas"] = extrair_qtd_parcelas(ws.cell(linha_encontrada, idx_parc).value)
        resultado["valor_parcela"] = formatar_moeda_br(ws.cell(linha_encontrada, idx_vparc).value)
        resultado["primeiro_vencimento"] = formatar_data_br(ws.cell(linha_encontrada, idx_dprim_venc).value)

        # 2. Localizar pasta do cliente no Dropbox (Busca Segura)
        pasta_cliente = self.obter_pasta_cliente_dropbox(nome_cliente, eh_juridico=eh_juridico)
        if pasta_cliente and pasta_cliente.exists():
            resultado["pasta_dropbox"] = str(pasta_cliente)
            arquivos_pasta = list(pasta_cliente.iterdir())

            # Sentença de Homologação
            for f in arquivos_pasta:
                if f.is_file() and f.suffix.lower() == ".pdf":
                    fn_norm = normalizar_texto(f.name)
                    if "SENTEN" in fn_norm or "HOMOLOGA" in fn_norm:
                        resultado["tem_sentenca"] = True
                        resultado["arquivo_sentenca"] = str(f.resolve())
                        resultado["documentos_encontrados"].append(f.name)
                        break

            # Termo / Ata de Audiência
            termo_pdf = None
            for f in arquivos_pasta:
                if f.is_file() and f.suffix.lower() == ".pdf":
                    fn_norm = normalizar_texto(f.name)
                    if any(k in fn_norm for k in ["TERMO", "AUDI", "ACORDO", "ATA"]) and not ("COMP" in fn_norm or "SENTEN" in fn_norm):
                        termo_pdf = f
                        resultado["documentos_encontrados"].append(f.name)
                        break

            # Cruzar com o texto do PDF
            if termo_pdf:
                try:
                    doc = pymupdf.open(termo_pdf)
                    texto_completo = ""
                    for p in doc:
                        texto_completo += p.get_text() + "\n"
                    doc.close()

                    if valor_acordo_str and valor_acordo_str in texto_completo:
                        resultado["validacao_pdf_notas"].append(f"Valor do acordo ({valor_acordo_str}) confirmado no PDF.")

                    # Valores do formulário de acordo: o TERMO manda; a planilha é reserva.
                    #   Valor inicial = valor original da causa ("valor atualizado totalizando R$ X")
                    #   Valor limite  = valor acordado ("pagamento do valor de R$ X" / "quantia de R$ X")
                    texto_1l = " ".join(texto_completo.split())
                    m_orig = re.search(r"valor atualizado totalizando\s*R\$\s*([\d.]+,\d{2})", texto_1l, re.I)
                    m_acordo = re.search(r"(?:pagamento do valor de|pagar a quantia de)\s*R\$\s*([\d.]+,\d{2})", texto_1l, re.I)
                    for chave, m, rotulo in (("valor_inicial", m_orig, "Valor inicial (original)"),
                                             ("valor_limite", m_acordo, "Valor limite (acordado)")):
                        if m:
                            v_termo = m.group(1)
                            if resultado[chave] and resultado[chave] != v_termo:
                                resultado["validacao_pdf_notas"].append(
                                    f"{rotulo}: termo R$ {v_termo} difere da planilha R$ {resultado[chave]}. Usando o do termo.")
                            resultado[chave] = v_termo
                            resultado[f"origem_{chave}"] = "termo"
                    if m_acordo:
                        resultado["valor_homologacao"] = resultado["valor_limite"]

                    # Valor da parcela do TERMO (usado nos boletos criados manualmente).
                    # Ex.: "mais 15 parcelas de R$ 335,39" / "24 parcelas de R$295,33"
                    m_parc = re.search(r"(\d+)\s*\(?[a-zç\s]*\)?\s*parcelas(?:\s+mensais)?,?\s*(?:no valor\s*)?de\s*R\$\s*([\d.]+,\d{2})",
                                       texto_1l, re.I)
                    if m_parc:
                        resultado["valor_parcela_termo"] = m_parc.group(2)
                        resultado["qtd_parcelas_termo"] = int(m_parc.group(1))

                    qtd = resultado["qtd_parcelas"]
                    if qtd > 0 and (f"{qtd} prestac" in normalizar_texto(texto_completo) or f"{qtd} parcel" in normalizar_texto(texto_completo) or f"{qtd}X" in normalizar_texto(texto_completo)):
                        resultado["validacao_pdf_notas"].append(f"Quantidade de parcelas ({qtd}) confirmada no PDF.")

                    vparc = resultado["valor_parcela"]
                    if vparc and vparc in texto_completo:
                        resultado["validacao_pdf_notas"].append(f"Valor da parcela ({vparc}) confirmado no PDF.")

                    if not resultado["processo"]:
                        proc_extraido = limpar_processo(texto_completo)
                        if proc_extraido:
                            resultado["processo"] = proc_extraido
                            resultado["validacao_pdf_notas"].append(f"Processo extraído do PDF: {proc_extraido}")
                except Exception as e_pdf:
                    resultado["validacao_pdf_notas"].append(f"Erro ao ler PDF do termo: {e_pdf}")
        else:
            resultado["validacao_pdf_notas"].append("Pasta do cliente não encontrada no Dropbox.")

        # Valor dos boletos criados manualmente: sempre o do termo; se o termo for escaneado
        # (sem texto legível), usa o da planilha.
        if resultado.get("valor_parcela_termo"):
            resultado["valor_parcela_boleto"] = resultado["valor_parcela_termo"]
            resultado["origem_valor_parcela"] = "termo"
        else:
            resultado["valor_parcela_boleto"] = resultado["valor_parcela"]
            resultado["origem_valor_parcela"] = "planilha (termo sem texto legível)"

        resultado["sucesso"] = True
        return resultado