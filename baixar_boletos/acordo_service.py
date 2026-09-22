"""
acordo_service.py
Serviço de automação Playwright para inclusão de Acordos Judiciais no Global Office.
Acessa o Cadastro de Devedores, abre o formulário de Acordo (ícone mãos dadas ),
preenche os dados validados (apenas PARCELADOS) e grava no sistema.
"""

import os
import sys
import time
from datetime import datetime
from pathlib import Path
from typing import Optional, Dict, Any, Callable

from playwright.sync_api import sync_playwright, Page, BrowserContext, Browser

from baixar_boletos.config import (
    GLOBAL_OFFICE_URL,
    GLOBAL_OFFICE_LOGIN,
    GLOBAL_OFFICE_USER,
    GLOBAL_OFFICE_PASS,
    STATE_FILE,
    USER_AGENT
)
from baixar_boletos.acordo_parser import AcordoParser, normalizar_texto

class AcordoService:
    def __init__(self, log_callback: Optional[Callable[[str], None]] = None, progress_callback: Optional[Callable] = None):
        self.log_callback = log_callback or (lambda msg: print(msg, flush=True))
        self.progress_callback = progress_callback
        self.playwright = None
        self.browser: Optional[Browser] = None
        self.context: Optional[BrowserContext] = None
        self.page: Optional[Page] = None
        self.parser = AcordoParser()
        self._cancelado = False

    def log(self, mensagem: str, nivel: str = "INFO"):
        hora = time.strftime("%H:%M:%S")
        texto = f"{hora} - {nivel} - {mensagem}"
        try:
            self.log_callback(texto)
        except UnicodeEncodeError:
            texto_limpo = texto.encode("ascii", "replace").decode("ascii")
            self.log_callback(texto_limpo)

    def cancelar(self):
        self._cancelado = True
        self.log("Cancelamento solicitado pelo usuário.", "WARNING")

    def _esperar_login_manual(self):
        self.log("Sessão não autenticada. Preenchendo credenciais...")
        browser = self.playwright.chromium.launch(headless=False)
        context = browser.new_context(user_agent=USER_AGENT)
        page = context.new_page()
        page.goto(GLOBAL_OFFICE_URL)

        if GLOBAL_OFFICE_LOGIN:
            try: page.locator("#EDTLOGIN").fill(GLOBAL_OFFICE_LOGIN)
            except Exception: pass
        if GLOBAL_OFFICE_USER:
            try: page.locator("#EDTUSUARIO").fill(GLOBAL_OFFICE_USER)
            except Exception: pass
        if GLOBAL_OFFICE_PASS:
            try: page.locator("#EDTSENHA").fill(GLOBAL_OFFICE_PASS)
            except Exception: pass

        self.log("Por favor, resolva o reCAPTCHA na janela aberta do navegador e clique em ENTRAR.")
        try:
            page.locator("#EDTLOGIN").wait_for(state="hidden", timeout=0)
            page.wait_for_timeout(3000)
            context.storage_state(path=str(STATE_FILE))
            self.log("Login efetuado com sucesso!")
        finally:
            browser.close()

    def _testar_sessao_ativa(self) -> bool:
        """Verifica se a sessão ativa ainda é reconhecida pelo servidor."""
        try:
            if self.page.locator("#EDTLOGIN").is_visible():
                return False
            has_seta = self.page.evaluate("typeof window.SetaAcao !== 'undefined'")
            if not has_seta:
                return False
            self.page.evaluate("if (typeof window.SetaAcao === 'function') window.SetaAcao('Menu', 'frmEmpresas');")
            self.page.wait_for_timeout(2500)
            linha = self.page.locator("tr:has-text('AC COBRANCA')").first
            if linha.is_visible(timeout=4000):
                return True
            return False
        except Exception:
            return False

    def iniciar_sessao(self, visivel: bool = True):
        """Inicia navegador."""
        self.log(f"Iniciando navegador (visível={1 if visivel else 0})...")
        self.playwright = sync_playwright().start()

        while True:
            state_file_path = Path(STATE_FILE)
            if not state_file_path.exists():
                local_state = Path("globaloffice_state.json")
                if local_state.exists():
                    state_file_path = local_state

            if not state_file_path.exists():
                self._esperar_login_manual()
                state_file_path = Path(STATE_FILE)

            launch_args = [
                "--start-maximized",
                "--disable-blink-features=AutomationControlled",
                "--window-size=1920,1080"
            ]

            self.browser = self.playwright.chromium.launch(
                headless=not visivel,
                args=launch_args
            )
            self.context = self.browser.new_context(
                storage_state=str(state_file_path),
                user_agent=USER_AGENT,
                viewport={"width": 1920, "height": 1080}
            )
            self.page = self.context.new_page()

            # Trata diálogos automáticos
            self.page.on("dialog", lambda dialog: dialog.accept())

            try:
                self.page.goto(GLOBAL_OFFICE_URL, timeout=60000)
                self.page.wait_for_timeout(3000)
            except Exception as e_nav:
                self.log(f"Falha de navegação na sessão salva: {e_nav}. Renovando login...", "WARNING")
                try: self.browser.close()
                except Exception: pass
                if state_file_path.exists():
                    try: state_file_path.unlink()
                    except Exception: pass
                continue

            if not self._testar_sessao_ativa():
                self.log("Sessão salva expirada no servidor. Solicitando novo login...", "WARNING")
                try: self.browser.close()
                except Exception: pass
                if state_file_path.exists():
                    try: state_file_path.unlink()
                    except Exception: pass
                self._esperar_login_manual()
                continue

            self.log("Sessão validada com sucesso.")
            self._selecionar_empresa()
            break

    def _selecionar_empresa(self):
        self.log("Confirmando seleção da empresa AC COBRANÇA...")
        try:
            linha = self.page.locator("tr:has-text('AC COBRANCA')").first
            if linha.is_visible(timeout=3000):
                botao_check = linha.locator("button[title*='principal'], a.btn-info, button.btn-info, button:has(.fa-check)").first
                if botao_check.is_visible():
                    botao_check.click()
                else:
                    self.page.evaluate("if(typeof window.SetaAcao === 'function') window.SetaAcao('Pri', '5');")
            else:
                self.page.evaluate("if(typeof window.SetaAcao === 'function') window.SetaAcao('Pri', '5');")

            self.page.wait_for_timeout(1500)

            # Confirma o popup SweetAlert: "Confirma marcar a empresa selecionada como Principal ?"
            btn_confirma_sweet = self.page.locator("button.confirm, .sweet-alert button:has-text('Confirmar'), button:has-text('Confirmar')").first
            if btn_confirma_sweet.is_visible(timeout=4000):
                self.log("Confirmando popup de empresa principal no sistema...")
                btn_confirma_sweet.click()
                self.page.wait_for_timeout(2500)

            # Garante que qualquer overlay residual seja removido
            try:
                self.page.evaluate("() => { const ov = document.querySelectorAll('.sweet-overlay, .sweet-alert'); ov.forEach(el => el.remove()); }")
            except Exception: pass

            self.log("Empresa AC COBRANÇA confirmada com sucesso.")
        except Exception as e:
            self.log(f"Aviso ao selecionar empresa: {e}", "WARNING")

    def acessar_cadastro_devedores(self):
        """Acessa a tela de Cadastro de devedores via SetaAcao e aguarda o grid."""
        self.log("Acessando tela de Cadastro de devedores...")
        try:
            self.page.evaluate("if (typeof window.SetaAcao === 'function') { window.SetaAcao('Menu', 'frmConsultaDevedoresExecucao'); }")
            self.page.wait_for_selector("#tabela_filter input", state="visible", timeout=35000)
            self.page.wait_for_timeout(1000)
            try:
                self.page.wait_for_selector('#tabela_processing', state='hidden', timeout=10000)
            except Exception: pass
            self.log("Tabela de Cadastro de devedores pronta.")
        except Exception as e_wait:
            self.log(f"Aviso ao aguardar tabela: {e_wait}", "WARNING")
            # Se timeout, tenta uma vez recarregar ou acionar via menu
            try:
                menu_jur = self.page.locator("b:has-text('Menu Jurídico'), span:has-text('Menu Jurídico')").first
                if menu_jur.is_visible():
                    menu_jur.click(force=True)
                    self.page.wait_for_timeout(1000)
                    self.page.locator("b:has-text('Cobranças/Execuções')").first.click(force=True)
                    self.page.wait_for_timeout(1000)
                    self.page.locator("em:has-text('Cadastro de devedores')").first.click(force=True)
                    self.page.wait_for_selector("#tabela_filter input", state="visible", timeout=20000)
            except Exception: pass

    def pesquisar_devedor_no_grid(self, nome_cliente: str) -> bool:
        """Localiza o devedor na tabela de devedores."""
        self.log(f"Pesquisando devedor no grid: '{nome_cliente}'...")
        try:
            campo_busca = self.page.locator('#tabela_filter input[type="search"], #tabela_filter input').first
            campo_busca.wait_for(state="visible", timeout=15000)
            campo_busca.click()
            self.page.keyboard.press("Control+A")
            self.page.keyboard.press("Backspace")
            self.page.wait_for_timeout(300)
            campo_busca.fill(nome_cliente)
            self.page.wait_for_timeout(2000)

            try:
                self.page.wait_for_selector("#tabela_processing", state="hidden", timeout=10000)
            except Exception: pass
            self.page.wait_for_timeout(1000)

            info_locator = self.page.locator("#tabela_info").first
            texto_info = info_locator.inner_text() if info_locator.is_visible() else ""
            if "0 até 0 de 0" in texto_info or "Mostrando 0" in texto_info:
                self.log(f"Nenhum registro encontrado no grid para '{nome_cliente}'.", "WARNING")
                return False

            linhas = self.page.locator("#tabela tbody tr").all()
            if not linhas:
                return False
            if "Nenhum registro" in linhas[0].inner_text():
                return False

            self.log(f"Devedor '{nome_cliente}' localizado no grid.")
            return True
        except Exception as e:
            self.log(f"Erro ao pesquisar devedor: {e}", "ERROR")
            return False

    def abrir_modal_acordo(self, nome_cliente: str = "") -> Optional[bool]:
        """Clica no botão de mãos dadas (fa-handshake-o) na linha do devedor. Retorna True se abriu, None se já tem acordo, False se erro."""
        self.log("Abrindo formulário de Acordo Judicial (ícone mãos dadas)...")
        try:
            linha = self.page.locator("#tabela tbody tr").first
            
            # Verifica se o devedor já possui acordo registrado no sistema
            html_linha = linha.evaluate("e => e.innerHTML")
            if "já tem acordo" in html_linha.lower():
                self.log(f"O devedor '{nome_cliente}' JÁ POSSUI ACORDO CADASTRADO no sistema (registro confirmado com sucesso).", "SUCCESS")
                return None

            btn_acordo = None

            # Tenta por seletores de ícone (fa-handshake-o, fa-handshake)
            for sel in [
                "button:has(i.fa-handshake-o)",
                "button:has(i.fa-handshake)",
                "button[title*='acordo' i]",
                "button[title*='Registrar' i]",
                "button[onclick*='acordo' i]",
            ]:
                try:
                    candidate = linha.locator(sel).first
                    if candidate.is_visible(timeout=500) and candidate.is_enabled():
                        btn_acordo = candidate
                        self.log(f"Botão de acordo localizado via seletor: {sel}")
                        break
                except Exception:
                    continue

            # Fallback: inspeciona HTML da linha para achar o botão correto
            if btn_acordo is None:
                try:
                    html_linha = linha.evaluate("e => e.innerHTML")
                    botoes = linha.locator("button").all()
                    self.log(f"Inspecionando {len(botoes)} botões na linha (fallback)...")
                    for i, btn in enumerate(botoes):
                        try:
                            btn_html = btn.evaluate("e => e.outerHTML").lower()
                            if any(k in btn_html for k in ["handshake", "acordo", "registr"]):
                                if btn.is_visible() and btn.is_enabled():
                                    btn_acordo = btn
                                    self.log(f"Botão de acordo encontrado no fallback (botão #{i})")
                                    break
                        except Exception:
                            continue
                except Exception:
                    pass

            if btn_acordo and btn_acordo.is_visible() and btn_acordo.is_enabled():
                btn_acordo.click()
                self.page.wait_for_timeout(2500)

                modal = self.page.locator("#Cadastro, form, div[role='dialog']").filter(has_text="Acordos Judiciais").first
                modal.wait_for(state="visible", timeout=15000)
                self.log("Formulário 'INCLUSÃO DE ACORDOS JUDICIAIS' aberto com sucesso!")
                return True
            else:
                # Log do HTML da linha para diagnóstico
                try:
                    html_diag = linha.evaluate("e => e.innerHTML")
                    self.log(f"HTML da linha (diagnóstico): {html_diag[:300]}", "WARNING")
                except Exception:
                    pass
                self.log("Botão de mãos dadas não localizado ou desabilitado na linha do devedor.", "ERROR")
                return False                return False

        except Exception as e:
            self.log(f"Erro ao abrir modal de acordos: {e}", "ERROR")
            return False

        except Exception as e:
            self.log(f"Erro ao abrir modal de acordos: {e}", "ERROR")
            return False

    def preencher_e_gravar_acordo(self, dados: Dict[str, Any]) -> bool:
        """
        Preenche todos os campos do formulário modal e salva no sistema.
        """
        modal = self.page.locator("#Cadastro, form, div[role='dialog']").filter(has_text="Acordos Judiciais").first

        try:
            # 1. BLOCO REGISTRO
            # Aguarda modal carregar completamente
            self.page.wait_for_timeout(1500)

            # Processo: tenta preencher com timeout curto.
            # Se o campo for readonly (qualquer motivo), simplesmente pula.
            input_proc = modal.locator("#EDTPROCESSO").first
            if input_proc.is_visible():
                val_proc = ""
                try:
                    val_proc = input_proc.input_value(timeout=3000).strip()
                except Exception:
                    pass

                if val_proc:
                    self.log(f"Processo já preenchido no modal: '{val_proc}'.")
                elif dados.get('processo'):
                    self.log(f"Tentando preencher Processo: {dados['processo']}")
                    try:
                        input_proc.fill(dados['processo'], timeout=3000)
                        self.log(f"Processo preenchido: {dados['processo']}")
                    except Exception:
                        # Campo readonly — normal para devedores já vinculados
                        self.log(
                            f"Campo EDTPROCESSO não editável (readonly). "
                            f"Sistema manterá o processo vinculado ao devedor."
                        )
            # Data de registro (garante data atual se vazia)
            input_dt_reg = modal.locator("#EDTDATAREGISTRO").first
            if input_dt_reg.is_visible() and not input_dt_reg.input_value().strip():
                input_dt_reg.fill(dados.get("data_registro", datetime.now().strftime("%d/%m/%Y")))

            # Valor Inicial (Dívida original / Valor da causa)
            if dados.get("valor_inicial"):
                self.log(f"Preenchendo Valor Inicial: R$ {dados['valor_inicial']}")
                input_vi = modal.locator("#EDTVALORINICIAL").first
                if input_vi.is_visible():
                    input_vi.click()
                    self.page.keyboard.press("Control+A")
                    self.page.keyboard.press("Backspace")
                    input_vi.fill(dados["valor_inicial"])

            # Valor Limite (Valor do acordo)
            if dados.get("valor_limite"):
                self.log(f"Preenchendo Valor Limite: R$ {dados['valor_limite']}")
                input_vl = modal.locator("#EDTVALORLIMITE").first
                if input_vl.is_visible():
                    input_vl.click()
                    self.page.keyboard.press("Control+A")
                    self.page.keyboard.press("Backspace")
                    input_vl.fill(dados["valor_limite"])

            # Vínculo do acordo: Titular
            radio_titular = modal.locator("#RDBTNTITULAR_INPUT, input[value='rdbtnTitular']").first
            if radio_titular.is_visible():
                radio_titular.check()

            # 2. BLOCO HOMOLOGAÇÃO
            # Arquivo Sentença
            if dados.get("tem_sentenca") and dados.get("arquivo_sentenca"):
                self.log(f"Anexando Sentença de Homologação: {Path(dados['arquivo_sentenca']).name}")
                modal.locator("#IWUPLOADER_INPUT").set_input_files(dados["arquivo_sentenca"])
            else:
                self.log("Sem Sentença de Homologação na pasta. Apagando campo Data da homologação...")
                self.page.evaluate("() => { const el = document.getElementById('EDTDATAHOMOLOGACAO'); if (el) { el.value = ''; el.dispatchEvent(new Event('change', {bubbles:true})); } }")

            # Valor da Homologação (Valor do acordo)
            if dados.get("valor_homologacao"):
                input_vh = modal.locator("#EDTVALORHOMOLOGADO").first
                if input_vh.is_visible():
                    input_vh.click()
                    self.page.keyboard.press("Control+A")
                    self.page.keyboard.press("Backspace")
                    input_vh.fill(dados["valor_homologacao"])

            # 3. BLOCO FINANCEIRO (Opções de pagamento)
            # Descrição do item de venda com validação na tabela flutuante
            self.log("Preenchendo Descrição do item de venda e validando na lista...")
            input_item = modal.locator("#EDTITEM").first
            if input_item.is_visible():
                input_item.click()
                input_item.fill("")
                input_item.press_sequentially("Prestação", delay=80)
                self.page.wait_for_timeout(1500)

                # Clica na linha da tabela flutuante LISTAITEM
                item_linha = self.page.locator(
                    "#LISTAITEM tr:has-text('Prestação'), #GRIDLISTAITEM tr:has-text('Prestação'), "
                    "td:has-text('Prestação de serviço')"
                ).first
                if item_linha.is_visible(timeout=5000):
                    item_linha.click()
                    self.page.wait_for_timeout(800)
                else:
                    self.page.keyboard.press("Enter")

                # Garante que a lista flutuante fechou para não cobrir outros campos
                self.page.evaluate("() => { const el = document.getElementById('LISTAITEM'); if (el) el.style.display = 'none'; }")
                self.page.wait_for_timeout(300)

            # Forma de pagamento, Modelo e Escritório
            self.log("Selecionando Forma de Pagamento (Boleto), Modelo (SICOOB- AC) e Escritório (Brasilia)...")
            self.page.evaluate("""() => {
                const f = document.getElementById('CBXFORMASDEPAGAMENTO');
                if (f) { f.value = '0'; f.dispatchEvent(new Event('change', { bubbles: true })); }
                const m = document.getElementById('CBXMODELO');
                if (m) { m.value = '0'; m.dispatchEvent(new Event('change', { bubbles: true })); }
                const e = document.getElementById('CBXESCRITORIOS');
                if (e) { e.value = '0'; e.dispatchEvent(new Event('change', { bubbles: true })); }
            }""")
            self.page.wait_for_timeout(800)

            # Entrada (Data e Valor)
            if dados.get("tem_entrada") and dados.get("valor_entrada"):
                dt_ent = dados['data_entrada']
                val_ent = dados['valor_entrada']
                self.log(f"Preenchendo Entrada: {dt_ent} - R$ {val_ent}")
                self.page.evaluate(f"""() => {{
                    const elDt = document.getElementById('EDTVENCIMENTODAENTRADA');
                    if (elDt) {{
                        elDt.value = '{dt_ent}';
                        elDt.dispatchEvent(new Event('input', {{ bubbles: true }}));
                        elDt.dispatchEvent(new Event('change', {{ bubbles: true }}));
                    }}
                    const elVal = document.getElementById('EDTVALORDAENTRADA');
                    if (elVal) {{
                        elVal.value = '{val_ent}';
                        elVal.dispatchEvent(new Event('input', {{ bubbles: true }}));
                        elVal.dispatchEvent(new Event('change', {{ bubbles: true }}));
                    }}
                }}""")
                self.page.keyboard.press("Escape")
            else:
                self.log("Acordo sem entrada (deixando campos de entrada vazios).")

            # Parcelamento 1 (Qtd.Parcelas, Valor das parcelas, 1º Vencimento)
            qtd1 = str(dados['qtd_parcelas'])
            val1 = dados['valor_parcela']
            venc1 = dados['primeiro_vencimento']
            self.log(f"Preenchendo Parcelamento 1: {qtd1}x de R$ {val1} (1º Vencimento: {venc1})")
            self.page.evaluate(f"""() => {{
                const elQtd = document.getElementById('EDTQTDPARCELAS1');
                if (elQtd) {{
                    elQtd.value = '{qtd1}';
                    elQtd.dispatchEvent(new Event('input', {{ bubbles: true }}));
                    elQtd.dispatchEvent(new Event('change', {{ bubbles: true }}));
                }}
                const elVal = document.getElementById('EDTVALORDASPARCELAS1');
                if (elVal) {{
                    elVal.value = '{val1}';
                    elVal.dispatchEvent(new Event('input', {{ bubbles: true }}));
                    elVal.dispatchEvent(new Event('change', {{ bubbles: true }}));
                }}
                const elVenc = document.getElementById('EDTPRIMEIROVENCIMENTO1');
                if (elVenc) {{
                    elVenc.value = '{venc1}';
                    elVenc.dispatchEvent(new Event('input', {{ bubbles: true }}));
                    elVenc.dispatchEvent(new Event('change', {{ bubbles: true }}));
                }}
            }}""")
            self.page.keyboard.press("Escape")
            self.page.wait_for_timeout(1000)

            # 4. GRAVAR DADOS
            self.log("Clicando em [GRAVAR DADOS]...")
            btn_gravar = modal.locator("button:has-text('Gravar Dados'), #BTNGRAVAR, .BTNGRAVARCSS").first
            if btn_gravar.is_visible():
                # Clica sem aguardar navegação (Global Office processa in-place, não navega)
                btn_gravar.click(no_wait_after=True)
                # Aguarda o sistema processar e exibir popups (pode demorar)
                self.page.wait_for_timeout(5000)

                # --- Popup 1: Confirmação de ausência de data de homologação ---
                # Usa locator direto para SweetAlert visível em vez de .all() que inclui hidden
                try:
                    sweet = self.page.locator(".sweet-alert").filter(has_text="homologa")
                    if sweet.is_visible(timeout=4000):
                        txt_h = sweet.inner_text()
                        self.log(f"Popup de homologação detectado: {txt_h[:80]}")
                        btn_sim = sweet.locator("button.confirm, button:has-text('Sim'), button:has-text('OK')").first
                        btn_sim.click(no_wait_after=True)
                        self.page.wait_for_timeout(6000)
                except Exception:
                    pass  # Sem popup de homologação, tudo bem

                # --- Popup 2: Confirmação de sucesso ---
                try:
                    sweet_suc = self.page.locator(".sweet-alert").first
                    if sweet_suc.is_visible(timeout=8000):
                        txt_s = sweet_suc.inner_text().strip()
                        self.log(f"Mensagem do sistema: {txt_s}")
                        btn_ok = sweet_suc.locator("button.confirm, button:has-text('OK'), button:has-text('Confirmar')").first
                        btn_ok.click(no_wait_after=True)
                        self.page.wait_for_timeout(2000)
                except Exception:
                    pass

                self.log(f"Acordo do cliente '{dados['nome_cliente']}' GRAVADO com sucesso!", "SUCCESS")
                return True
            else:
                self.log("Botão [GRAVAR DADOS] não localizado no modal.", "ERROR")
                return False

        except Exception as e:
            self.log(f"Erro ao preencher/gravar acordo: {e}", "ERROR")
            return False

        def subir_acordo_completo(self, nome_cliente: str, visivel: bool = True) -> bool:
        """Executa o ciclo completo de subir o acordo para um cliente."""
        self.log("=" * 60)
        self.log(f"INICIANDO ENVIO DO ACORDO JUDICIAL: {nome_cliente}")
        self.log("=" * 60)

        # 1. Ler e validar dados do cliente
        dados = self.parser.carregar_dados_cliente(nome_cliente)
        if not dados["sucesso"]:
            self.log(f"Não é possível prosseguir: {dados['motivo_erro']}", "WARNING")
            return False

        if not dados["eh_parcelado"]:
            self.log(f"Cliente {nome_cliente} ignorado: APENAS ACORDOS PARCELADOS são permitidos.", "WARNING")
            return False

        self.log(f"Dados validados: Processo {dados['processo']} | Causa: R$ {dados['valor_inicial']} | Acordo: R$ {dados['valor_limite']}")
        self.log(f"Condições: {dados['qtd_parcelas']}x de R$ {dados['valor_parcela']} | Entrada: {dados['valor_entrada'] or 'Nenhuma'}")
        if dados["validacao_pdf_notas"]:
            for nota in dados["validacao_pdf_notas"]:
                self.log(f"Validação Documental: {nota}")

        # 2. Sessão e navegação
        if not self.page:
            self.iniciar_sessao(visivel=visivel)

        self.acessar_cadastro_devedores()

        # 3. Pesquisa e abertura do modal
        achou = self.pesquisar_devedor_no_grid(nome_cliente)
        if not achou:
            self.log(f"Devedor '{nome_cliente}' não encontrado no grid de devedores.", "ERROR")
            return False

        abriu = self.abrir_modal_acordo(nome_cliente)
        if abriu is None:
            # Já tem acordo cadastrado
            return True
        if not abriu:
            return False

        # 4. Preenchimento e gravação
        sucesso = self.preencher_e_gravar_acordo(dados)
        return sucesso

    def lista_parcelados_planilha(self):
        """Retorna lista de nomes de clientes PARCELADOS da planilha, na ordem da planilha."""
        try:
            import openpyxl
            from baixar_boletos.acordo_parser import normalizar_texto
        except ImportError:
            import openpyxl
            from acordo_parser import normalizar_texto

        parser = self.parser
        caminho = parser.caminho_planilha
        if not caminho.exists():
            self.log(f"Planilha não encontrada: {caminho}", "ERROR")
            return []

        wb = openpyxl.load_workbook(str(caminho), data_only=True)
        ws = wb.active
        for s in wb.sheetnames:
            if "FECHADO" in s.upper() or "ENCONTRADO" in s.upper():
                ws = wb[s]
                break

        header = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        col_map = {}
        for idx, h in enumerate(header, 1):
            if h:
                col_map[normalizar_texto(str(h))] = idx

        idx_nome = col_map.get("BENEFICIARIO") or 3
        idx_cond = col_map.get("CONDICOES DO ACORDO") or 7

        parcelados = []
        for r in range(2, ws.max_row + 1):
            nome = ws.cell(r, idx_nome).value
            cond = str(ws.cell(r, idx_cond).value or "").strip().upper()
            if nome and "PARCEL" in cond:
                parcelados.append(str(nome).strip())

        self.log(f"Planilha lida: {len(parcelados)} clientes PARCELADOS encontrados.")
        return parcelados

    def fechar_sessao(self):
        try:
            if self.context: self.context.close()
            if self.browser: self.browser.close()
            if self.playwright: self.playwright.stop()
        except Exception:
            pass
        self.log("Sessão finalizada.")
