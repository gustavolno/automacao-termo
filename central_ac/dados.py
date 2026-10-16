"""Leitura dos dados de clientes: planilha da Jaci (acordos + operador) e conciliação do Sicoob
(adimplência). Tudo somente leitura. A interface pública (Repositorio) é o que o PostgreSQL vai
substituir: as telas e a API não sabem de onde os dados vêm."""
import re
import threading
from datetime import date, datetime
from pathlib import Path
from typing import Optional

import openpyxl

from baixar_boletos.acordo_parser import AcordoParser, extrair_qtd_parcelas
from baixar_boletos.boletos_faltantes import cronograma_parcelas
from baixar_boletos.config import normalizar_texto

from . import settings, store

STATUS = {"ADIMPLENTE": "ADIMPLENTE", "INADIMPLENTE": "INADIMPLENTE",
          "SEM REGISTRO – VERIFICAR": "SEM_REGISTRO", "EM ANÁLISE / DIVERGENTE": "EM_ANALISE"}


def so_digitos(v) -> str:
    return re.sub(r"\D", "", str(v or ""))


def _cpf(v) -> str:
    d = so_digitos(v)
    return d.zfill(11) if 8 <= len(d) <= 11 else d


def _data(v) -> Optional[date]:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()[:10]
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def _moeda(v) -> Optional[float]:
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v or "").replace("R$", "").strip()
    if not s:
        return None
    s = s.replace(".", "").replace(",", ".") if "," in s else s
    try:
        return float(s)
    except ValueError:
        return None


def _iso(d: Optional[date]) -> Optional[str]:
    return d.isoformat() if d else None


class Repositorio:
    def __init__(self):
        self._lock = threading.Lock()
        self._chave = None
        self._clientes: dict = {}
        self._parcelas: dict = {}
        self._fontes: dict = {}

    # ── carga (recarrega sozinho quando um dos arquivos muda) ──
    def _arquivos(self):
        parser = AcordoParser()
        plan = Path(parser.caminho_planilha)
        concs = sorted(settings.RAIZ.glob("CONCILIACAO_ADIMPLENCIA*.xlsx"), key=lambda p: p.stat().st_mtime)
        conc = concs[-1] if concs else None
        return parser, plan, conc

    def _garantir(self):
        with self._lock:
            parser, plan, conc = self._arquivos()
            chave = (str(plan), plan.stat().st_mtime if plan.exists() else 0,
                     str(conc), conc.stat().st_mtime if conc else 0)
            if chave != self._chave:
                self._carregar(parser, plan, conc)
                self._chave = chave

    def _carregar(self, parser, plan: Path, conc: Optional[Path]):
        concil_cpf, concil_nome, parcelas, vencido_aberto = {}, {}, {}, {}
        if conc:
            wb = openpyxl.load_workbook(conc, data_only=True, read_only=True)
            for r in wb["Consolidado"].iter_rows(min_row=2, values_only=True):
                if not r[0]:
                    continue
                chave_p = _cpf(r[1]) or normalizar_texto(str(r[0]))
                st_p = str(r[7] or "")        # soma por status da PARCELA (é como o relatório de conciliação valora)
                vencido_aberto[(chave_p, st_p)] = vencido_aberto.get((chave_p, st_p), 0.0) + (_moeda(r[4]) or 0.0)
                parcelas.setdefault(chave_p, []).append({
                    "mes": str(r[2] or ""), "vencimento": _iso(_data(r[3])), "devido": _moeda(r[4]),
                    "pago": _moeda(r[5]), "data_pagamento": _iso(_data(r[6])),
                    "status": str(r[7] or ""), "obs": str(r[8] or "")})
            for r in wb["Clientes"].iter_rows(min_row=2, values_only=True):
                if not r[0]:
                    continue
                status = STATUS.get(r[2], "EM_ANALISE")
                chave_c = _cpf(r[1]) or normalizar_texto(str(r[0]))     # a MESMA chave da aba Consolidado
                # Valor em aberto como no relatório: inadimplente = parcelas inadimplentes; sem registro = parcelas sem
                # registro; em análise = o valor da aba Clientes (a aba soma também parcelas ainda a vencer).
                par = {"INADIMPLENTE": "INADIMPLENTE", "SEM_REGISTRO": "SEM REGISTRO – VERIFICAR"}.get(status)
                valor = round(vencido_aberto.get((chave_c, par), 0.0), 2) if par else (_moeda(r[5]) or 0.0)
                d = {"status": status, "parcelas_vencidas": int(_moeda(r[3]) or 0),
                     "parcelas_pendentes": int(_moeda(r[4]) or 0), "valor_pendente": valor,
                     "pendente_desde": _data(r[6]), "pago_ate": str(r[7] or ""), "motivo": str(r[8] or ""), "cpf_conc": chave_c}
                concil_cpf[_cpf(r[1])] = d
                concil_nome[normalizar_texto(str(r[0]))] = d
            wb.close()

        ws = parser._obter_worksheet()
        cab = {normalizar_texto(str(c.value)): i for i, c in enumerate(ws[1]) if c.value}
        col = lambda nome, padrao=None: cab.get(nome, padrao)  # noqa: E731
        i_nome, i_cpf = col("BENEFICIARIO", 3), col("CPF", 4)
        campos = {
            "processo": col("PROCESSO"), "fechado_em": col("ACORDO FECHADO EM"), "condicao": col("CONDICOES DO ACORDO"),
            "valor_acordo": col("VALOR INTEGRAL DO ACORDO"), "data_entrada": col("DATA DA ENTRADA /QUITACAO"),
            "valor_entrada": col("VALOR DA ENTRADA/ QUITACAO"), "qtd_parcelas": col("PARCELAS"),
            "valor_parcela": col("VALOR DA PARCELA"), "primeiro_venc": col("PAGAMENTO INICIAL"),
            "operador": col("OPERADOR RESPONSAVEL"), "envio": col("ENVIO DE BOLETOS"),
        }
        clientes, usados, valor_contado = {}, set(), set()
        for row in ws.iter_rows(min_row=2, values_only=True):
            nome = str(row[i_nome] or "").strip() if i_nome is not None else ""
            if not nome:
                continue
            cpf = _cpf(row[i_cpf]) if i_cpf is not None else ""
            cid = cpf if len(cpf) == 11 else "n-" + re.sub(r"\W+", "-", normalizar_texto(nome)).lower()
            while cid in usados:
                cid += "-2"
            usados.add(cid)
            v = {k: (row[i] if i is not None else None) for k, i in campos.items()}
            condicao = normalizar_texto(str(v["condicao"] or ""))
            parcelado = "PARCEL" in condicao
            op = normalizar_texto(str(v["operador"] or ""))
            qtd = extrair_qtd_parcelas(v["qtd_parcelas"])
            primeiro = _data(v["primeiro_venc"])
            original = concil_cpf.get(cpf) or concil_nome.get(normalizar_texto(nome))
            repetido = False
            if original:
                repetido = original["cpf_conc"] in valor_contado     # mesma pessoa em outra linha da planilha
                valor_contado.add(original["cpf_conc"])
                conc_c = dict(original)
                if repetido:                                 # a dívida não entra duas vezes nos totais
                    conc_c["valor_pendente"] = 0.0
            else:
                conc_c = {"status": "SEM_CONCILIACAO" if parcelado else "NAO_PARCELADO", "parcelas_vencidas": 0,
                          "parcelas_pendentes": 0, "valor_pendente": 0.0, "pendente_desde": None, "pago_ate": "", "motivo": "", "cpf_conc": ""}
            venc_futuros = []
            if parcelado and primeiro and qtd:
                venc_futuros = [d for d in cronograma_parcelas(primeiro, qtd) if d >= date.today()]
            clientes[cid] = {
                "id": cid, "nome": nome, "cpf": cpf, "operador": op,
                "condicao": str(v["condicao"] or "").strip(), "parcelado": parcelado, "processo": str(v["processo"] or ""),
                "fechado_em": str(v["fechado_em"] or ""), "valor_acordo": _moeda(v["valor_acordo"]),
                "valor_entrada": _moeda(v["valor_entrada"]), "data_entrada": _iso(_data(v["data_entrada"])),
                "qtd_parcelas": qtd, "valor_parcela": _moeda(v["valor_parcela"]), "primeiro_venc": _iso(primeiro),
                "proximo_venc": _iso(venc_futuros[0]) if venc_futuros else None,
                "envio_boletos": normalizar_texto(str(v["envio"] or "")), "repetido": repetido,
                **{k: (_iso(x) if isinstance(x, date) else x) for k, x in conc_c.items()},
            }
        self._clientes, self._parcelas = clientes, parcelas
        self._fontes = {"planilha": plan.name, "planilha_em": datetime.fromtimestamp(plan.stat().st_mtime).strftime("%d/%m/%Y %H:%M"),
                        "conciliacao": conc.name if conc else None,
                        "conciliacao_em": datetime.fromtimestamp(conc.stat().st_mtime).strftime("%d/%m/%Y") if conc else None}

    # ── consultas ──
    def _carteira(self, c: dict, atrib: dict) -> str:
        if c["id"] in atrib:
            return atrib[c["id"]]
        return c["operador"] if c["operador"] in settings.NEGOCIADORAS else settings.A_DISTRIBUIR

    def clientes(self, carteira: Optional[str] = None, status: Optional[str] = None, q: str = "") -> list:
        self._garantir()
        atrib, contatos = store.atribuicoes(), store.resumo_contatos()
        alvo = normalizar_texto(q)
        hoje = date.today()
        out = []
        for c in self._clientes.values():
            cart = self._carteira(c, atrib)
            if carteira and carteira != settings.TODAS and cart != carteira:
                continue
            if status and status != "TODOS":
                if status == "LEMBRETES":
                    lem = (contatos.get(c["id"]) or {}).get("lembrar_em")
                    if not lem or lem > hoje.isoformat():
                        continue
                elif c["status"] != status:
                    continue
            if alvo and alvo not in normalizar_texto(c["nome"]) and alvo not in c["cpf"]:
                continue
            d = dict(c, carteira=cart, carteira_nome=settings.NEGOCIADORAS.get(cart, "A distribuir" if cart == settings.A_DISTRIBUIR else cart))
            pd = _data(d["pendente_desde"])
            d["dias_atraso"] = (hoje - pd).days if pd and d["status"] == "INADIMPLENTE" else 0
            d.update({"ultimo_contato": (contatos.get(c["id"]) or {}).get("ultimo"),
                      "lembrar_em": (contatos.get(c["id"]) or {}).get("lembrar_em")})
            out.append(d)
        return out

    def cliente(self, cid: str) -> Optional[dict]:
        self._garantir()
        c = self._clientes.get(cid)
        if not c:
            return None
        return next(x for x in self.clientes(settings.TODAS) if x["id"] == cid)

    def parcelas(self, cid: str) -> list:
        self._garantir()
        c = self._clientes.get(cid)
        if not c:
            return []
        linhas = (self._parcelas.get(c["cpf"]) or self._parcelas.get(c.get("cpf_conc", ""))
                  or self._parcelas.get(normalizar_texto(c["nome"])) or [])
        return sorted(linhas, key=lambda p: p["vencimento"] or "")

    def carteiras(self) -> list:
        todos = self.clientes(settings.TODAS)
        out = []
        for chave, nome in list(settings.NEGOCIADORAS.items()) + [(settings.A_DISTRIBUIR, "A distribuir"), (settings.TODAS, "Visão da gestão")]:
            meus = todos if chave == settings.TODAS else [c for c in todos if c["carteira"] == chave]
            out.append({"id": chave, "nome": nome, "clientes": len(meus),
                        "inadimplentes": sum(c["status"] == "INADIMPLENTE" for c in meus)})
        return out

    def resumo(self, carteira: str = settings.TODAS) -> dict:
        meus = self.clientes(carteira)
        por = {}
        for c in meus:
            s = por.setdefault(c["status"], {"clientes": 0, "valor": 0.0})
            s["clientes"] += 1
            s["valor"] += c["valor_pendente"] or 0.0
        return {"carteira": carteira, "total": len(meus), "por_status": por,
                "lembretes_hoje": sum(1 for c in meus if c["lembrar_em"] and c["lembrar_em"] <= date.today().isoformat())}

    def fontes(self) -> dict:
        self._garantir()
        return self._fontes


repo = Repositorio()
