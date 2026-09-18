import os
import sys
import re
import time
from pathlib import Path
from datetime import datetime
import openpyxl
from dotenv import load_dotenv
from playwright.sync_api import sync_playwright

# Carrega variáveis de ambiente
load_dotenv()

GLOBAL_OFFICE_URL = os.getenv("GLOBAL_OFFICE_URL", "https://www.globalofficeweb.com.br/")
GLOBAL_OFFICE_LOGIN = os.getenv("GLOBAL_OFFICE_LOGIN", "")
GLOBAL_OFFICE_USER = os.getenv("GLOBAL_OFFICE_USER", "")
GLOBAL_OFFICE_PASS = os.getenv("GLOBAL_OFFICE_PASS", "")

STATE_FILE = Path(r"baixar_boletos\globaloffice_state.json")
USER_AGENT = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"

class VerificadorDevedoresService:
    def __init__(self, log_callback=None):
        self.log_callback = log_callback or (lambda msg: print(msg))
        self.playwright = None
        self.browser = None
        self.context = None
        self.page = None
        self._cancelado = False

    def log(self, mensagem: str, nivel: str = "INFO"):
        hora = time.strftime("%H:%M:%S")
        texto = f"{hora} - {nivel} - {mensagem}"
        self.log_callback(texto)

    def _esperar_login_manual(self):
        self.log("Sessão não autenticada. Preenchendo credenciais...")
        browser = self.playwright.chromium.launch(headless=False)
        context = browser.new_context(user_agent=USER_AGENT)
        page = context.new_page()
        page.goto(GLOBAL_OFFICE_URL)

        if GLOBAL_OFFICE_LOGIN:
            try: page.locator("#EDTLOGIN").fill(GLOBAL_OFFICE_LOGIN)
            except Exception: pass
        if GLOBAL_OFFICE_USER:
            try: page.locator("#EDTUSUARIO").fill(GLOBAL_OFFICE_USER)
            except Exception: pass
        if GLOBAL_OFFICE_PASS:
            try: page.locator("#EDTSENHA").fill(GLOBAL_OFFICE_PASS)
            except Exception: pass

        self.log("Por favor, resolva o reCAPTCHA na janela aberta do navegador e clique em ENTRAR.")
        try:
            page.locator("#EDTLOGIN").wait_for(state="hidden", timeout=0)
            page.wait_for_timeout(3000)
            context.storage_state(path=str(STATE_FILE))
            self.log("Login efetuado com sucesso!")
            self.log(f"Sessão salva com sucesso em '{STATE_FILE}'.")
        finally:
            browser.close()

    def iniciar_sessao(self, visivel: bool = False):
        self.log(f"Iniciando navegador (visivel={1 if visivel else 0})...")
        self.playwright = sync_playwright().start()

        while True:
            if not STATE_FILE.exists():
                self._esperar_login_manual()

            launch_args = [
                "--start-maximized",
                "--disable-blink-features=AutomationControlled",
                "--window-size=1920,1080"
            ]

            self.browser = self.playwright.chromium.launch(
                headless=not visivel,
                args=launch_args
            )
            self.context = self.browser.new_context(
                storage_state=str(STATE_FILE),
                user_agent=USER_AGENT,
                viewport={"width": 1920, "height": 1080}
            )
            self.page = self.context.new_page()
            self.page.on("dialog", lambda dialog: dialog.accept())

            try:
                self.page.goto(GLOBAL_OFFICE_URL, timeout=60000)
                self.page.wait_for_timeout(3000)
            except Exception as e_nav:
                self.log(f"Falha de conexão na sessão salva: {e_nav}. Renovando login...", "WARNING")
                if self.context: self.context.close()
                if self.browser: self.browser.close()
                STATE_FILE.unlink(missing_ok=True)
                continue

            is_login_visible = False
            try:
                is_login_visible = self.page.locator("#EDTLOGIN").is_visible()
            except Exception: pass

            has_seta = False
            try:
                has_seta = self.page.evaluate("typeof window.SetaAcao !== 'undefined'")
            except Exception: pass

            if is_login_visible or not has_seta:
                self.log("Sessão expirada. Solicitando novo login...", "WARNING")
                if self.context: self.context.close()
                if self.browser: self.browser.close()
                STATE_FILE.unlink(missing_ok=True)
                continue
            else:
                self.log("Sessão validada com sucesso.")
                self._selecionar_empresa()
                break

    def _selecionar_empresa(self):
        self.log("Selecionando a empresa AC COBRANÇA...")
        try:
            self.page.evaluate("if (typeof window.SetaAcao === 'function') window.SetaAcao('Menu', 'frmEmpresas');")
            self.page.wait_for_timeout(2500)

            try:
                linha = self.page.locator("tr:has-text('AC COBRANCA')").first
                if linha.is_visible(timeout=4000):
                    botao_check = linha.locator("button[title*='principal'], a.btn-info, button.btn-info, button:has(.fa-check)").first
                    if botao_check.is_visible():
                        botao_check.click()
                    else:
                        self.page.evaluate("if(typeof window.SetaAcao === 'function') window.SetaAcao('Pri', '5');")
                else:
                    self.page.evaluate("if(typeof window.SetaAcao === 'function') window.SetaAcao('Pri', '5');")
            except Exception:
                self.page.evaluate("if(typeof window.SetaAcao === 'function') window.SetaAcao('Pri', '5');")

            self.page.wait_for_timeout(2500)
        except Exception as e:
            self.log(f"Aviso ao selecionar empresa: {e}", "WARNING")

    def acessar_cadastro_devedores(self):
        self.log("Acessando tela de Cadastro de devedores...")
        
        # Tenta abrir pelo menu lateral
        try:
            menu_jur = self.page.locator("a:has-text('Menu Jurídico'), span:has-text('Menu Jurídico')").first
            if menu_jur.is_visible():
                menu_jur.click()
                self.page.wait_for_timeout(1000)

            cobrancas = self.page.locator("a:has-text('Cobranças/Execuções'), span:has-text('Cobranças/Execuções'), em:has-text('Cobranças/Execuções')").first
            if cobrancas.is_visible():
                cobrancas.click()
                self.page.wait_for_timeout(1000)

            link_devedores = self.page.locator("a:has-text('Cadastro de devedores'), em:has-text('Cadastro de devedores')").first
            if link_devedores.is_visible():
                link_devedores.click()
                self.page.wait_for_timeout(3000)
        except Exception as e_menu:
            self.log(f"Tentativa pelo menu lateral: {e_menu}", "WARNING")

        self.log("Aguardando carregamento da tabela de devedores...")
        try:
            self.page.wait_for_selector('table[id*="tabela"], #tabela, table.dataTable', state='visible', timeout=30000)
            self.page.wait_for_timeout(2000)
            self.log("Tabela de Cadastro de devedores pronta.")
        except Exception as e_wait:
            self.log(f"Aviso ao aguardar tabela: {e_wait}", "WARNING")

    def verificar_devedor_no_grid(self, nome_devedor: str) -> bool:
        """
        Pesquisa o devedor no campo de pesquisa do grid e verifica se há registro.
        Retorna True se encontrado, False caso contrário.
        """
        nome_limpo = re.sub(r'\s+', ' ', nome_devedor).strip().upper()
        if not nome_limpo:
            return False

        # Localiza o campo de pesquisa do DataTables
        campo_pesquisa = self.page.locator("input[type='search'], #tabela_filter input, input[aria-controls*='tabela']").first
        if not campo_pesquisa.is_visible():
            campo_pesquisa = self.page.locator("input.form-control-sm, input[placeholder*='Pesquisar']").first

        try:
            campo_pesquisa.click()
            campo_pesquisa.fill("")
            self.page.wait_for_timeout(300)
            campo_pesquisa.fill(nome_limpo)
            self.page.wait_for_timeout(1500)

            # Aguarda a resposta do grid
            try:
                self.page.wait_for_selector('#tabela_processing, .dataTables_processing', state='hidden', timeout=10000)
            except Exception:
                pass
            self.page.wait_for_timeout(1000)

            # Verifica texto informativo de paginação / contagem
            info_locator = self.page.locator("#tabela_info, .dataTables_info").first
            texto_info = info_locator.inner_text() if info_locator.is_visible() else ""

            if "0 até 0 de 0" in texto_info or "Mostrando 0" in texto_info:
                return False

            # Verifica linhas do corpo da tabela
            linhas = self.page.locator("table[id*='tabela'] tbody tr, #tabela tbody tr").all()
            if not linhas:
                return False

            texto_primeira_linha = linhas[0].inner_text().strip()
            if "Nenhum registro" in texto_primeira_linha or "Aguarde" in texto_primeira_linha or "Não foram encontrados" in texto_primeira_linha:
                return False

            return True

        except Exception as e:
            self.log(f"Erro ao pesquisar '{nome_limpo}': {e}", "ERROR")
            return False

    def fechar(self):
        try:
            if self.context: self.context.close()
            if self.browser: self.browser.close()
            if self.playwright: self.playwright.stop()
        except Exception:
            pass


def carregar_dados_planilha(caminho_planilha: str, nome_aba: str = "FECHADOS "):
    wb = openpyxl.load_workbook(caminho_planilha, data_only=True)
    if nome_aba not in wb.sheetnames:
        candidatas = [s for s in wb.sheetnames if 'FECHADO' in s.upper() or 'JACI' in s.upper()]
        nome_aba = candidatas[0] if candidatas else wb.sheetnames[0]

    ws = wb[nome_aba]
    print(f"Lendo aba '{nome_aba}'...")

    cabecalho = []
    idx_beneficiario = -1
    for c in range(1, ws.max_column + 1):
        v = ws.cell(1, c).value
        texto_col = str(v).strip().upper() if v else f"COLUNA_{c}"
        cabecalho.append(texto_col)
        if "BENEFICI" in texto_col or "NOME" in texto_col or "CLIENTE" in texto_col or "DEVEDOR" in texto_col:
            if idx_beneficiario == -1:
                idx_beneficiario = c - 1

    if idx_beneficiario == -1:
        idx_beneficiario = 2

    registros = []
    for r in range(2, ws.max_row + 1):
        linha_vals = [ws.cell(r, c).value for c in range(1, len(cabecalho) + 1)]
        nome = str(linha_vals[idx_beneficiario] or "").strip()
        if nome and nome != "None":
            dados_linha = {}
            for i, col in enumerate(cabecalho):
                dados_linha[col] = str(linha_vals[i] if i < len(linha_vals) and linha_vals[i] is not None else "").strip()
            
            registros.append({
                "linha": r,
                "nome": nome,
                "dados": dados_linha
            })

    return cabecalho, registros


def main():
    print("=" * 70)
    print("ROBÔ DE VERIFICAÇÃO DE DEVEDORES NO GLOBAL OFFICE")
    print("=" * 70)

    caminho_planilha = Path("planilha_jaci.xlsx")
    if not caminho_planilha.exists():
        print(f"[ERRO] Planilha '{caminho_planilha}' não encontrada na raiz.")
        sys.exit(1)

    cabecalho, registros = carregar_dados_planilha(str(caminho_planilha), nome_aba="FECHADOS ")
    total = len(registros)
    print(f"Total de devedores carregados da planilha: {total}")

    if total == 0:
        print("[AVISO] Nenhum registro encontrado na planilha.")
        sys.exit(0)

    service = VerificadorDevedoresService()
    
    arq_nao_encontrados = Path("devedores_nao_encontrados.txt")
    arq_encontrados_txt = Path("devedores_encontrados.txt")
    arq_encontrados_xlsx = Path("devedores_encontrados.xlsx")

    with open(arq_nao_encontrados, "w", encoding="utf-8") as f:
        f.write(f"=== DEVEDORES NÃO ENCONTRADOS NO CADASTRO DE DEVEDORES ===\n")
        f.write(f"Data de Execução: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}\n")
        f.write(f"Origem: {caminho_planilha.name} (Aba FECHADOS)\n")
        f.write("=" * 70 + "\n\n")

    with open(arq_encontrados_txt, "w", encoding="utf-8") as f:
        f.write(f"=== DEVEDORES ENCONTRADOS NO CADASTRO DE DEVEDORES ===\n")
        f.write(f"Data de Execução: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}\n")
        f.write(f"Origem: {caminho_planilha.name} (Aba FECHADOS)\n")
        f.write("=" * 70 + "\n\n")

    wb_saida = openpyxl.Workbook()
    ws_saida = wb_saida.active
    ws_saida.title = "ENCONTRADOS"
    ws_saida.append(cabecalho + ["STATUS_GLOBAL_OFFICE", "HORA_VERIFICACAO"])

    encontrados_count = 0
    nao_encontrados_count = 0

    try:
        visivel = "--visivel" in sys.argv or "-v" in sys.argv
        service.iniciar_sessao(visivel=visivel)
        service.acessar_cadastro_devedores()

        print("\nIniciando varredura no grid...")
        for idx, item in enumerate(registros, 1):
            nome = item["nome"]
            dados = item["dados"]

            print(f"[{idx}/{total}] Pesquisando: {nome}...", end=" ", flush=True)
            achou = service.verificar_devedor_no_grid(nome)

            hora_atual = datetime.now().strftime("%H:%M:%S")

            if achou:
                print(">> ENCONTRADO! <<")
                encontrados_count += 1
                
                with open(arq_encontrados_txt, "a", encoding="utf-8") as f:
                    f.write(f"[{encontrados_count}] {nome}\n")
                    for k, v in dados.items():
                        if v:
                            f.write(f"    - {k}: {v}\n")
                    f.write("\n")

                linha_excel = [dados.get(col, "") for col in cabecalho]
                linha_excel.append("ENCONTRADO")
                linha_excel.append(hora_atual)
                ws_saida.append(linha_excel)
                wb_saida.save(arq_encontrados_xlsx)

            else:
                print(">> NÃO ENCONTRADO <<")
                nao_encontrados_count += 1
                
                with open(arq_nao_encontrados, "a", encoding="utf-8") as f:
                    cpf = dados.get("CPF", "")
                    processo = dados.get("PROCESSO", "")
                    uf = dados.get("UF", "")
                    f.write(f"[{nao_encontrados_count}] {nome} | CPF: {cpf} | UF: {uf} | Processo: {processo}\n")

    except KeyboardInterrupt:
        print("\n\n[AVISO] Processamento interrompido pelo usuário.")
    except Exception as e:
        print(f"\n\n[ERRO CRÍTICO] {e}")
    finally:
        service.fechar()
        wb_saida.save(arq_encontrados_xlsx)

    print("\n" + "=" * 70)
    print("RELATÓRIO FINAL DA VERIFICAÇÃO:")
    print(f"  - Total processado: {encontrados_count + nao_encontrados_count} de {total}")
    print(f"  - Encontrados no sistema: {encontrados_count}")
    print(f"  - NÃO encontrados no sistema: {nao_encontrados_count}")
    print("=" * 70)
    print(f"Arquivos gerados com sucesso:")
    print(f"  1. {arq_nao_encontrados.resolve()}")
    print(f"  2. {arq_encontrados_txt.resolve()}")
    print(f"  3. {arq_encontrados_xlsx.resolve()}")
    print("=" * 70)

if __name__ == "__main__":
    main()
